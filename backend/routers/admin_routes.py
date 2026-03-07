from fastapi import APIRouter, HTTPException, Depends, UploadFile, File
from typing import Optional
import csv
import io
import uuid
from datetime import datetime, timezone

from database import db
from auth import get_current_user, require_admin, require_admin_or_manager
from models import VehicleCreate, Vehicle
from services import log_audit, recalculate_product_stock

router = APIRouter()


# ============= VEHICLES =============
@router.get("/vehicles")
async def get_vehicles(current_user: dict = Depends(get_current_user)):
    vehicles = await db.vehicles.find({}, {"_id": 0}).sort("name", 1).to_list(100)
    return vehicles


@router.post("/vehicles")
async def create_vehicle(data: VehicleCreate, admin: dict = Depends(require_admin_or_manager)):
    doc = {"id": str(uuid.uuid4()), **data.model_dump(), "created_at": datetime.now(timezone.utc).isoformat()}
    await db.vehicles.insert_one(doc)
    await log_audit(admin, "create", "vehicle", doc["id"], {"name": data.name, "plate": data.plate_number})
    return {k: v for k, v in doc.items() if k != '_id'}


@router.put("/vehicles/{vehicle_id}")
async def update_vehicle(vehicle_id: str, data: VehicleCreate, admin: dict = Depends(require_admin_or_manager)):
    result = await db.vehicles.find_one_and_update(
        {"id": vehicle_id}, {"$set": data.model_dump()}, return_document=True, projection={"_id": 0}
    )
    if not result:
        raise HTTPException(status_code=404, detail="Armada tidak ditemukan")
    await log_audit(admin, "update", "vehicle", vehicle_id, data.model_dump())
    return result


@router.delete("/vehicles/{vehicle_id}")
async def delete_vehicle(vehicle_id: str, admin: dict = Depends(require_admin_or_manager)):
    result = await db.vehicles.delete_one({"id": vehicle_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Armada tidak ditemukan")
    await db.partner_stores.update_many({"vehicle_id": vehicle_id}, {"$set": {"vehicle_id": ""}})
    await log_audit(admin, "delete", "vehicle", vehicle_id, {})
    return {"message": "Armada berhasil dihapus"}


# ============= AUDIT LOG =============
@router.get("/audit-log")
async def get_audit_log(entity_type: Optional[str] = None, limit: int = 100, admin: dict = Depends(require_admin_or_manager)):
    query = {}
    if entity_type:
        query["entity_type"] = entity_type
    logs = await db.audit_log.find(query, {"_id": 0}).sort("timestamp", -1).to_list(limit)
    return logs


# ============= DATA IMPORT =============
IMPORT_TYPES_INFO = [
    {"key": "warehouses", "label": "Gudang", "order": 1, "columns": ["nama", "unit", "urutan"]},
    {"key": "categories", "label": "Kategori", "order": 2, "columns": ["nama", "deskripsi"]},
    {"key": "suppliers", "label": "Supplier", "order": 3, "columns": ["nama", "kontak", "telepon", "email", "alamat"]},
    {"key": "products", "label": "Produk", "order": 4, "columns": ["nama", "barcode", "kategori", "supplier", "harga_jual", "harga_beli", "harga_box", "isi_per_box", "min_stok", "kadaluarsa", "batch", "deskripsi"]},
    {"key": "warehouse_stock", "label": "Stok per Gudang", "order": 5, "columns": ["produk", "gudang", "stok"]},
    {"key": "purchases", "label": "Pembelian Supplier", "order": 6, "columns": ["no_pembelian", "supplier", "tanggal", "produk", "jumlah", "harga_satuan", "total_pembelian", "dibayar", "sisa", "gudang", "catatan"]},
    {"key": "payment_history", "label": "Pembayaran Hutang Supplier", "order": 7, "columns": ["supplier", "tanggal", "jumlah", "catatan"]},
    {"key": "partner_stores", "label": "Kios Mitra", "order": 8, "columns": ["nama", "pemilik", "telepon", "alamat", "limit_kredit", "hutang_saat_ini"]},
    {"key": "partner_orders", "label": "Pesanan Kios Mitra", "order": 9, "columns": ["no_pesanan", "kios_mitra", "tanggal", "produk", "jumlah", "harga", "gudang", "diskon", "total_pesanan", "dibayar", "sisa", "sales", "status"]},
]


async def parse_upload(file: UploadFile):
    content = await file.read()
    fname = (file.filename or "").lower()
    if fname.endswith('.xlsx') or fname.endswith('.xls'):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
        headers = []
        for cell in ws[1]:
            if cell.value is not None:
                headers.append(str(cell.value).strip().lower())
            else:
                break
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            rd = {}
            for idx, val in enumerate(row):
                if idx < len(headers):
                    rd[headers[idx]] = val if val is not None else ""
            if any(str(v).strip() != "" for v in rd.values()):
                rows.append(rd)
        return rows
    else:
        text = content.decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(text))
        return [{k.strip().lower(): (v.strip() if isinstance(v, str) else v) for k, v in row.items() if k} for row in reader]


def safe_float(v, default=0):
    try:
        return float(str(v).replace(",", ".").replace(" ", "")) if v != "" and v is not None else default
    except Exception:
        return default


def safe_int(v, default=0):
    try:
        return int(float(str(v).replace(",", ".").replace(" ", ""))) if v != "" and v is not None else default
    except Exception:
        return default


def safe_str(v, default=""):
    return str(v).strip() if v is not None and str(v).strip() else default


@router.get("/import/types")
async def get_import_types(admin: dict = Depends(require_admin)):
    return IMPORT_TYPES_INFO


@router.get("/import/counts")
async def get_import_counts(admin: dict = Depends(require_admin)):
    collections = ["warehouses", "categories", "suppliers", "products", "warehouse_stock", "purchases", "payment_history", "partner_stores", "partner_orders"]
    result = {}
    for c in collections:
        result[c] = await db[c].count_documents({})
    return result


@router.get("/import/template/{import_type}")
async def get_import_template(import_type: str, admin: dict = Depends(require_admin)):
    from starlette.responses import StreamingResponse
    info = next((t for t in IMPORT_TYPES_INFO if t["key"] == import_type), None)
    if not info:
        raise HTTPException(status_code=404, detail="Tipe import tidak ditemukan")
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(info["columns"])
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]), media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=template_{import_type}.csv"}
    )


@router.post("/import/{import_type}")
async def import_data(import_type: str, file: UploadFile = File(...), admin: dict = Depends(require_admin)):
    try:
        rows = await parse_upload(file)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Gagal membaca file: {str(e)}")
    if not rows:
        raise HTTPException(status_code=400, detail="File kosong atau format tidak sesuai")

    now = datetime.now(timezone.utc)
    results = {"success": 0, "skipped": 0, "errors": []}

    if import_type == "warehouses":
        existing = {w["name"].lower(): w for w in await db.warehouses.find({}, {"_id": 0}).to_list(100)}
        max_order = max((w.get("sort_order", 0) for w in existing.values()), default=0)
        for i, row in enumerate(rows, 1):
            nama = safe_str(row.get("nama"))
            if not nama:
                results["errors"].append(f"Baris {i}: Nama kosong")
                continue
            if nama.lower() in existing:
                results["skipped"] += 1
                continue
            max_order += 1
            doc = {"id": nama.lower().replace(" ", "_").replace("-", "_"), "name": nama, "unit": safe_str(row.get("unit"), "dus").lower(), "sort_order": safe_int(row.get("urutan"), max_order)}
            if doc["unit"] not in ("dus", "satuan"):
                doc["unit"] = "dus"
            await db.warehouses.insert_one(doc)
            existing[nama.lower()] = doc
            results["success"] += 1

    elif import_type == "categories":
        existing = {c["name"].lower() for c in await db.categories.find({}, {"_id": 0, "name": 1}).to_list(10000)}
        for i, row in enumerate(rows, 1):
            nama = safe_str(row.get("nama"))
            if not nama:
                results["errors"].append(f"Baris {i}: Nama kosong")
                continue
            if nama.lower() in existing:
                results["skipped"] += 1
                continue
            doc = {"id": str(uuid.uuid4()), "name": nama, "description": safe_str(row.get("deskripsi")), "created_at": now.isoformat()}
            await db.categories.insert_one(doc)
            existing.add(nama.lower())
            results["success"] += 1

    elif import_type == "suppliers":
        existing = {s["name"].lower() for s in await db.suppliers.find({}, {"_id": 0, "name": 1}).to_list(10000)}
        for i, row in enumerate(rows, 1):
            nama = safe_str(row.get("nama"))
            if not nama:
                results["errors"].append(f"Baris {i}: Nama kosong")
                continue
            if nama.lower() in existing:
                results["skipped"] += 1
                continue
            doc = {"id": str(uuid.uuid4()), "name": nama, "contact_person": safe_str(row.get("kontak")), "phone": safe_str(row.get("telepon")), "email": safe_str(row.get("email")), "address": safe_str(row.get("alamat")), "created_at": now.isoformat()}
            await db.suppliers.insert_one(doc)
            existing.add(nama.lower())
            results["success"] += 1

    elif import_type == "products":
        cats = {c["name"].lower(): c["id"] for c in await db.categories.find({}, {"_id": 0}).to_list(10000)}
        sups = {s["name"].lower(): s["id"] for s in await db.suppliers.find({}, {"_id": 0}).to_list(10000)}
        existing = {p["name"].lower() for p in await db.products.find({}, {"_id": 0, "name": 1}).to_list(100000)}
        for i, row in enumerate(rows, 1):
            nama = safe_str(row.get("nama"))
            if not nama:
                results["errors"].append(f"Baris {i}: Nama kosong")
                continue
            if nama.lower() in existing:
                results["skipped"] += 1
                continue
            cat_name = safe_str(row.get("kategori")).lower()
            sup_name = safe_str(row.get("supplier")).lower()
            cat_id = cats.get(cat_name, "")
            sup_id = sups.get(sup_name, "")
            if cat_name and not cat_id:
                results["errors"].append(f"Baris {i}: Kategori '{row.get('kategori')}' tidak ditemukan")
                continue
            if sup_name and not sup_id:
                results["errors"].append(f"Baris {i}: Supplier '{row.get('supplier')}' tidak ditemukan")
                continue
            doc = {
                "id": str(uuid.uuid4()), "name": nama, "barcode": safe_str(row.get("barcode")),
                "category_id": cat_id, "supplier_id": sup_id, "stock": 0,
                "price": safe_float(row.get("harga_jual")), "cost_price": safe_float(row.get("harga_beli")),
                "box_price": safe_float(row.get("harga_box")), "items_per_box": max(1, safe_int(row.get("isi_per_box"), 1)),
                "min_stock_alert": safe_int(row.get("min_stok"), 10), "expiry_date": safe_str(row.get("kadaluarsa")),
                "batch_number": safe_str(row.get("batch")), "description": safe_str(row.get("deskripsi")),
                "image": "", "created_at": now.isoformat(), "updated_at": now.isoformat()
            }
            await db.products.insert_one(doc)
            existing.add(nama.lower())
            results["success"] += 1

    elif import_type == "warehouse_stock":
        prods = {p["name"].lower(): p for p in await db.products.find({}, {"_id": 0}).to_list(100000)}
        whs = {w["name"].lower(): w["id"] for w in await db.warehouses.find({}, {"_id": 0}).to_list(100)}
        affected = set()
        for i, row in enumerate(rows, 1):
            prod_name = safe_str(row.get("produk")).lower()
            wh_name = safe_str(row.get("gudang")).lower()
            stok = safe_int(row.get("stok"))
            if not prod_name or not wh_name:
                results["errors"].append(f"Baris {i}: Produk atau gudang kosong")
                continue
            prod = prods.get(prod_name)
            wh_id = whs.get(wh_name)
            if not prod:
                results["errors"].append(f"Baris {i}: Produk '{row.get('produk')}' tidak ditemukan")
                continue
            if not wh_id:
                results["errors"].append(f"Baris {i}: Gudang '{row.get('gudang')}' tidak ditemukan")
                continue
            await db.warehouse_stock.update_one({"product_id": prod["id"], "warehouse_id": wh_id}, {"$set": {"stock": stok}}, upsert=True)
            affected.add(prod["id"])
            results["success"] += 1
        for pid in affected:
            await recalculate_product_stock(pid)

    elif import_type == "purchases":
        sups = {s["name"].lower(): s for s in await db.suppliers.find({}, {"_id": 0}).to_list(10000)}
        prods = {p["name"].lower(): p for p in await db.products.find({}, {"_id": 0}).to_list(100000)}
        whs = {w["name"].lower(): w["id"] for w in await db.warehouses.find({}, {"_id": 0}).to_list(100)}
        groups = {}
        for i, row in enumerate(rows, 1):
            key = safe_str(row.get("no_pembelian")) or f"auto_{i}"
            if key not in groups:
                groups[key] = {"rows": [], "indices": []}
            groups[key]["rows"].append(row)
            groups[key]["indices"].append(i)
        for key, group in groups.items():
            first = group["rows"][0]
            sup_name = safe_str(first.get("supplier")).lower()
            sup = sups.get(sup_name)
            if not sup:
                for idx in group["indices"]:
                    results["errors"].append(f"Baris {idx}: Supplier '{first.get('supplier')}' tidak ditemukan")
                continue
            items = []
            has_error = False
            for j, row in enumerate(group["rows"]):
                prod_name = safe_str(row.get("produk")).lower()
                if not prod_name:
                    continue
                prod = prods.get(prod_name)
                if not prod:
                    results["errors"].append(f"Baris {group['indices'][j]}: Produk '{row.get('produk')}' tidak ditemukan")
                    has_error = True
                    continue
                jumlah = safe_int(row.get("jumlah"))
                harga = safe_float(row.get("harga_satuan"))
                items.append({"product_id": prod["id"], "product_name": prod["name"], "quantity": jumlah, "cost_price": harga, "subtotal": jumlah * harga})
            if has_error or not items:
                continue
            total = safe_float(first.get("total_pembelian")) or sum(it["subtotal"] for it in items)
            dibayar = safe_float(first.get("dibayar"))
            sisa = safe_float(first.get("sisa")) if safe_str(first.get("sisa")) else total - dibayar
            tanggal = safe_str(first.get("tanggal")) or now.strftime("%Y-%m-%d")
            gudang_id = whs.get(safe_str(first.get("gudang")).lower(), "gudang_utama")
            doc = {
                "id": str(uuid.uuid4()), "supplier_id": sup["id"], "supplier_name": sup["name"],
                "items": items, "total": total, "paid": dibayar, "remaining": sisa,
                "warehouse_id": gudang_id, "notes": safe_str(first.get("catatan")),
                "created_at": tanggal + "T00:00:00+00:00", "date": tanggal
            }
            await db.purchases.insert_one(doc)
            results["success"] += 1

    elif import_type == "payment_history":
        sups = {s["name"].lower(): s for s in await db.suppliers.find({}, {"_id": 0}).to_list(10000)}
        for i, row in enumerate(rows, 1):
            sup_name = safe_str(row.get("supplier")).lower()
            sup = sups.get(sup_name)
            if not sup:
                results["errors"].append(f"Baris {i}: Supplier '{row.get('supplier')}' tidak ditemukan")
                continue
            jumlah = safe_float(row.get("jumlah"))
            if jumlah <= 0:
                results["errors"].append(f"Baris {i}: Jumlah harus > 0")
                continue
            tanggal = safe_str(row.get("tanggal")) or now.strftime("%Y-%m-%d")
            doc = {
                "id": str(uuid.uuid4()), "type": "supplier", "reference_id": sup["id"],
                "supplier_id": sup["id"], "supplier_name": sup["name"], "amount": jumlah,
                "notes": safe_str(row.get("catatan")) or f"Import: Pembayaran ke {sup['name']}",
                "created_at": tanggal + "T00:00:00+00:00", "date": tanggal
            }
            await db.payment_history.insert_one(doc)
            results["success"] += 1

    elif import_type == "partner_stores":
        existing = {s["name"].lower() for s in await db.partner_stores.find({}, {"_id": 0, "name": 1}).to_list(10000)}
        for i, row in enumerate(rows, 1):
            nama = safe_str(row.get("nama"))
            if not nama:
                results["errors"].append(f"Baris {i}: Nama kosong")
                continue
            if nama.lower() in existing:
                results["skipped"] += 1
                continue
            doc = {
                "id": str(uuid.uuid4()), "name": nama, "owner_name": safe_str(row.get("pemilik")),
                "phone": safe_str(row.get("telepon")), "address": safe_str(row.get("alamat")),
                "latitude": None, "longitude": None,
                "credit_limit": safe_float(row.get("limit_kredit")),
                "current_debt": safe_float(row.get("hutang_saat_ini")),
                "created_at": now.isoformat()
            }
            await db.partner_stores.insert_one(doc)
            existing.add(nama.lower())
            results["success"] += 1

    elif import_type == "partner_orders":
        stores = {s["name"].lower(): s for s in await db.partner_stores.find({}, {"_id": 0}).to_list(10000)}
        prods = {p["name"].lower(): p for p in await db.products.find({}, {"_id": 0}).to_list(100000)}
        whs = {w["name"].lower(): w["id"] for w in await db.warehouses.find({}, {"_id": 0}).to_list(100)}
        canvassers_map = {c["name"].lower(): c for c in await db.canvassers.find({}, {"_id": 0}).to_list(1000)}
        groups = {}
        for i, row in enumerate(rows, 1):
            key = safe_str(row.get("no_pesanan")) or f"auto_{i}"
            if key not in groups:
                groups[key] = {"rows": [], "indices": []}
            groups[key]["rows"].append(row)
            groups[key]["indices"].append(i)
        for key, group in groups.items():
            first = group["rows"][0]
            store_name = safe_str(first.get("kios_mitra")).lower()
            store = stores.get(store_name)
            if not store:
                for idx in group["indices"]:
                    results["errors"].append(f"Baris {idx}: Kios '{first.get('kios_mitra')}' tidak ditemukan")
                continue
            items = []
            has_error = False
            for j, row in enumerate(group["rows"]):
                prod_name = safe_str(row.get("produk")).lower()
                if not prod_name:
                    continue
                prod = prods.get(prod_name)
                if not prod:
                    results["errors"].append(f"Baris {group['indices'][j]}: Produk '{row.get('produk')}' tidak ditemukan")
                    has_error = True
                    continue
                jumlah = safe_int(row.get("jumlah"))
                harga = safe_float(row.get("harga"))
                gudang_id = whs.get(safe_str(row.get("gudang")).lower(), "")
                items.append({"product_id": prod["id"], "product_name": prod["name"], "quantity": jumlah, "price": harga, "subtotal": jumlah * harga, "warehouse_id": gudang_id, "sell_unit": ""})
            if has_error or not items:
                continue
            subtotal = sum(it["subtotal"] for it in items)
            diskon = safe_float(first.get("diskon"))
            total = safe_float(first.get("total_pesanan")) or (subtotal - diskon)
            dibayar = safe_float(first.get("dibayar"))
            sisa = safe_float(first.get("sisa")) if safe_str(first.get("sisa")) else total - dibayar
            tanggal = safe_str(first.get("tanggal")) or now.strftime("%Y-%m-%d")
            sales_name = safe_str(first.get("sales")).lower()
            canvasser = canvassers_map.get(sales_name)
            order_number = key if not key.startswith("auto_") else f"PO-IMP-{str(uuid.uuid4())[:8].upper()}"
            doc = {
                "id": str(uuid.uuid4()), "order_number": order_number,
                "partner_store_id": store["id"], "partner_store_name": store["name"],
                "items": items, "subtotal": subtotal, "discount": diskon,
                "total": total, "paid": dibayar, "remaining": sisa,
                "delivery_date": tanggal,
                "canvasser_id": canvasser["id"] if canvasser else "",
                "canvasser_name": canvasser["name"] if canvasser else safe_str(first.get("sales")),
                "warehouse_id": whs.get(safe_str(first.get("gudang")).lower(), ""),
                "status": safe_str(first.get("status")) or "completed",
                "created_at": tanggal + "T00:00:00+00:00", "date": tanggal, "has_return": False
            }
            await db.partner_orders.insert_one(doc)
            results["success"] += 1
    else:
        raise HTTPException(status_code=400, detail=f"Tipe import '{import_type}' tidak didukung")

    return {"type": import_type, "total_rows": len(rows), "success": results["success"], "skipped": results["skipped"], "error_count": len(results["errors"]), "errors": results["errors"][:50]}


@router.delete("/import/clear/{collection_type}")
async def clear_collection(collection_type: str, admin: dict = Depends(require_admin)):
    allowed = ["warehouses", "categories", "suppliers", "products", "warehouse_stock", "purchases", "payment_history", "partner_stores", "partner_orders"]
    if collection_type not in allowed:
        raise HTTPException(status_code=400, detail="Tipe tidak didukung")
    count = await db[collection_type].count_documents({})
    await db[collection_type].delete_many({})
    return {"message": f"{count} data dihapus dari {collection_type}"}
