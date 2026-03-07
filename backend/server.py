from fastapi import FastAPI, APIRouter, HTTPException, Depends, status, UploadFile, File, Request, Query
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict, EmailStr
from typing import List, Optional, Dict, Any
import uuid
from datetime import datetime, timezone, timedelta
from passlib.context import CryptContext
from jose import JWTError, jwt
import csv
import io

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")
security = HTTPBearer()

SECRET_KEY = os.environ.get('SECRET_KEY', 'your-secret-key-change-this-in-production')
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_DAYS = 7

def hash_password(password: str) -> str:
    return pwd_context.hash(password)

def verify_password(plain_password: str, hashed_password: str) -> bool:
    return pwd_context.verify(plain_password, hashed_password)

def create_access_token(data: dict):
    to_encode = data.copy()
    expire = datetime.now(timezone.utc) + timedelta(days=ACCESS_TOKEN_EXPIRE_DAYS)
    to_encode.update({"exp": expire})
    return jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)

async def get_current_user(credentials: HTTPAuthorizationCredentials = Depends(security)):
    try:
        token = credentials.credentials
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        user_id: str = payload.get("sub")
        if user_id is None:
            raise HTTPException(status_code=401, detail="Invalid token")
    except JWTError:
        raise HTTPException(status_code=401, detail="Invalid token")
    
    user = await db.users.find_one({"id": user_id}, {"_id": 0})
    if user is None:
        raise HTTPException(status_code=401, detail="User not found")
    return user

async def require_admin(current_user: dict = Depends(get_current_user)):
    if current_user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin access required")
    return current_user

async def require_admin_or_manager(current_user: dict = Depends(get_current_user)):
    if current_user.get("role") not in ("admin", "manager"):
        raise HTTPException(status_code=403, detail="Admin or Manager access required")
    return current_user

# Role permissions
ROLE_PERMISSIONS = {
    "admin": ["dashboard", "products", "categories", "suppliers", "supplier_debts", "partner_stores", "partner_orders", "canvassers", "analytics", "reports", "users", "store_settings", "import_data", "printer", "profile", "vehicles", "audit_log", "pos"],
    "manager": ["dashboard", "products", "categories", "suppliers", "supplier_debts", "partner_stores", "partner_orders", "canvassers", "analytics", "reports", "store_settings", "import_data", "printer", "profile", "vehicles", "audit_log", "pos"],
    "office_staff": ["dashboard", "products", "categories", "suppliers", "supplier_debts", "partner_stores", "partner_orders", "analytics", "reports", "profile", "vehicles", "pos"],
    "cashier": ["pos", "profile"],
    "canvasser": ["canvasser_dashboard", "profile"],
}

class UserCreate(BaseModel):
    email: EmailStr
    password: str
    name: str
    role: str = "cashier"
    canvasser_id: Optional[str] = None

class UserLogin(BaseModel):
    email: EmailStr
    password: str

class User(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    email: str
    name: str
    role: str
    canvasser_id: Optional[str] = None
    created_at: str

class TokenResponse(BaseModel):
    token: str
    user: User

class CategoryCreate(BaseModel):
    name: str
    description: Optional[str] = ""

class Category(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    name: str
    description: str
    created_at: str

class SupplierCreate(BaseModel):
    name: str
    contact_person: Optional[str] = ""
    phone: Optional[str] = ""
    email: Optional[str] = ""
    address: Optional[str] = ""

class Supplier(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    name: str
    contact_person: str
    phone: str
    email: str
    address: str
    created_at: str

class ProductCreate(BaseModel):
    name: str
    barcode: Optional[str] = ""
    category_id: Optional[str] = ""
    supplier_id: Optional[str] = ""
    stock: int = 0
    warehouse_id: Optional[str] = "gudang_utama"
    price: float
    cost_price: float
    box_price: Optional[float] = 0
    items_per_box: int = 1
    initial_paid: Optional[float] = 0
    expiry_date: Optional[str] = ""
    batch_number: Optional[str] = ""
    description: Optional[str] = ""
    image: Optional[str] = ""
    min_stock_alert: int = 10

class Product(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    name: str
    barcode: str
    category_id: str
    supplier_id: str
    stock: int
    price: float
    cost_price: float
    box_price: float = 0
    items_per_box: int = 1
    expiry_date: str
    batch_number: str
    description: str
    image: str
    min_stock_alert: int
    created_at: str
    updated_at: str

class TransactionItem(BaseModel):
    product_id: str
    product_name: str
    quantity: int
    price: float
    subtotal: float
    warehouse_id: Optional[str] = ""
    sell_unit: Optional[str] = ""

class TransactionCreate(BaseModel):
    items: List[TransactionItem]
    subtotal: float
    tax: float = 0
    discount: float = 0
    total: float
    payment_method: str
    payment_amount: float
    change: float

class Transaction(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    transaction_number: str
    cashier_id: str
    cashier_name: str
    items: List[TransactionItem]
    subtotal: float
    tax: float
    discount: float
    total: float
    payment_method: str
    payment_amount: float
    change: float
    created_at: str
    date: str

@api_router.post("/auth/register", response_model=User)
async def register(user_data: UserCreate, admin: dict = Depends(require_admin)):
    existing = await db.users.find_one({"email": user_data.email}, {"_id": 0})
    if existing:
        raise HTTPException(status_code=400, detail="Email already registered")
    
    user_doc = {
        "id": str(uuid.uuid4()),
        "email": user_data.email,
        "password": hash_password(user_data.password),
        "name": user_data.name,
        "role": user_data.role,
        "canvasser_id": user_data.canvasser_id,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    
    await db.users.insert_one(user_doc)
    user_doc.pop("password")
    return user_doc

@api_router.post("/auth/login", response_model=TokenResponse)
async def login(credentials: UserLogin):
    user = await db.users.find_one({"email": credentials.email}, {"_id": 0})
    if not user or not verify_password(credentials.password, user["password"]):
        raise HTTPException(status_code=401, detail="Invalid email or password")
    
    token = create_access_token({"sub": user["id"]})
    user.pop("password")
    return {"token": token, "user": user}

@api_router.get("/auth/me", response_model=User)
async def get_me(current_user: dict = Depends(get_current_user)):
    return current_user

class ProfileUpdate(BaseModel):
    name: Optional[str] = None
    email: Optional[EmailStr] = None
    current_password: Optional[str] = None
    new_password: Optional[str] = None

class StoreSettings(BaseModel):
    model_config = ConfigDict(extra="ignore")
    store_name: str
    address: str
    phone: str
    email: str
    logo_url: Optional[str] = ""

class PurchaseCreate(BaseModel):
    supplier_id: str
    items: List[Dict[str, Any]]
    total: float
    paid: float = 0
    notes: Optional[str] = ""
    warehouse_id: Optional[str] = "gudang_utama"

class Purchase(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    supplier_id: str
    supplier_name: str
    items: List[Dict[str, Any]]
    total: float
    paid: float
    remaining: float
    notes: str
    created_at: str
    date: str

class PartnerStoreCreate(BaseModel):
    name: str
    owner_name: str
    phone: str
    address: str
    latitude: Optional[float] = None
    longitude: Optional[float] = None
    credit_limit: float = 0
    vehicle_id: Optional[str] = ""

class PartnerStore(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    name: str
    owner_name: str
    phone: str
    address: str
    latitude: Optional[float] = None
    longitude: Optional[float] = None
    credit_limit: float
    current_debt: float
    vehicle_id: Optional[str] = ""
    created_at: str

class PartnerOrderCreate(BaseModel):
    partner_store_id: str
    items: List[TransactionItem]
    subtotal: float
    discount: float = 0
    total: float
    paid: float = 0
    delivery_date: Optional[str] = ""
    canvasser_id: Optional[str] = ""
    warehouse_id: Optional[str] = ""

class PartnerOrder(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    order_number: str
    partner_store_id: str
    partner_store_name: str
    items: List[TransactionItem]
    subtotal: float
    discount: float
    total: float
    paid: float
    remaining: float
    delivery_date: str
    canvasser_id: str
    canvasser_name: str
    status: str
    created_at: str
    date: str
    has_return: Optional[bool] = False

class CanvasserCreate(BaseModel):
    name: str
    phone: str
    vehicle_number: str
    vehicle_type: str = "Mobil Box"

class Canvasser(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    name: str
    phone: str
    vehicle_number: str
    vehicle_type: str
    is_active: bool
    created_at: str

@api_router.put("/auth/profile", response_model=User)
async def update_profile(profile_data: ProfileUpdate, current_user: dict = Depends(get_current_user)):
    update_fields = {}
    
    if profile_data.name:
        update_fields["name"] = profile_data.name
    
    if profile_data.email:
        existing = await db.users.find_one({"email": profile_data.email, "id": {"$ne": current_user["id"]}}, {"_id": 0})
        if existing:
            raise HTTPException(status_code=400, detail="Email already in use")
        update_fields["email"] = profile_data.email
    
    if profile_data.new_password:
        if not profile_data.current_password:
            raise HTTPException(status_code=400, detail="Current password required")
        
        user_full = await db.users.find_one({"id": current_user["id"]})
        if not verify_password(profile_data.current_password, user_full["password"]):
            raise HTTPException(status_code=400, detail="Current password incorrect")
        
        update_fields["password"] = hash_password(profile_data.new_password)
    
    if not update_fields:
        raise HTTPException(status_code=400, detail="No fields to update")
    
    result = await db.users.find_one_and_update(
        {"id": current_user["id"]},
        {"$set": update_fields},
        return_document=True,
        projection={"_id": 0, "password": 0}
    )
    
    return result

@api_router.get("/categories", response_model=List[Category])
async def get_categories(current_user: dict = Depends(get_current_user)):
    categories = await db.categories.find({}, {"_id": 0}).sort("name", 1).to_list(1000)
    return categories

@api_router.post("/categories", response_model=Category)
async def create_category(data: CategoryCreate, admin: dict = Depends(require_admin)):
    category_doc = {
        "id": str(uuid.uuid4()),
        "name": data.name,
        "description": data.description,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.categories.insert_one(category_doc)
    return category_doc

@api_router.put("/categories/{category_id}", response_model=Category)
async def update_category(category_id: str, data: CategoryCreate, admin: dict = Depends(require_admin)):
    result = await db.categories.find_one_and_update(
        {"id": category_id},
        {"$set": {"name": data.name, "description": data.description}},
        return_document=True,
        projection={"_id": 0}
    )
    if not result:
        raise HTTPException(status_code=404, detail="Category not found")
    return result

@api_router.delete("/categories/{category_id}")
async def delete_category(category_id: str, admin: dict = Depends(require_admin)):
    result = await db.categories.delete_one({"id": category_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Category not found")
    return {"message": "Category deleted"}

@api_router.get("/suppliers")
async def get_suppliers(
    search: Optional[str] = None,
    page: Optional[int] = None,
    per_page: int = 15,
    current_user: dict = Depends(get_current_user)
):
    query = {}
    if search:
        query["name"] = {"$regex": search, "$options": "i"}
    if page is not None:
        total = await db.suppliers.count_documents(query)
        skip = (page - 1) * per_page
        suppliers = await db.suppliers.find(query, {"_id": 0}).sort("name", 1).skip(skip).limit(per_page).to_list(per_page)
        return {"data": suppliers, "total": total, "page": page, "per_page": per_page, "total_pages": (total + per_page - 1) // per_page}
    suppliers = await db.suppliers.find(query, {"_id": 0}).sort("name", 1).to_list(10000)
    return suppliers

@api_router.post("/suppliers", response_model=Supplier)
async def create_supplier(data: SupplierCreate, admin: dict = Depends(require_admin)):
    supplier_doc = {
        "id": str(uuid.uuid4()),
        "name": data.name,
        "contact_person": data.contact_person,
        "phone": data.phone,
        "email": data.email,
        "address": data.address,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.suppliers.insert_one(supplier_doc)
    return supplier_doc

@api_router.put("/suppliers/{supplier_id}", response_model=Supplier)
async def update_supplier(supplier_id: str, data: SupplierCreate, admin: dict = Depends(require_admin)):
    result = await db.suppliers.find_one_and_update(
        {"id": supplier_id},
        {"$set": data.model_dump()},
        return_document=True,
        projection={"_id": 0}
    )
    if not result:
        raise HTTPException(status_code=404, detail="Supplier not found")
    return result

@api_router.delete("/suppliers/{supplier_id}")
async def delete_supplier(supplier_id: str, admin: dict = Depends(require_admin)):
    result = await db.suppliers.delete_one({"id": supplier_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Supplier not found")
    return {"message": "Supplier deleted"}

@api_router.get("/products")
async def get_products(
    search: Optional[str] = None,
    category_id: Optional[str] = None,
    page: Optional[int] = None,
    per_page: int = 15,
    current_user: dict = Depends(get_current_user)
):
    query = {}
    if search:
        query["$or"] = [
            {"name": {"$regex": search, "$options": "i"}},
            {"barcode": {"$regex": search, "$options": "i"}}
        ]
    if category_id:
        query["category_id"] = category_id
    
    if page is not None:
        total = await db.products.count_documents(query)
        skip = (page - 1) * per_page
        products = await db.products.find(query, {"_id": 0}).sort("name", 1).skip(skip).limit(per_page).to_list(per_page)
        return {"data": products, "total": total, "page": page, "per_page": per_page, "total_pages": (total + per_page - 1) // per_page}
    
    products = await db.products.find(query, {"_id": 0}).sort("name", 1).to_list(100000)
    return products

@api_router.get("/products/barcode/{barcode}", response_model=Product)
async def get_product_by_barcode(barcode: str, current_user: dict = Depends(get_current_user)):
    product = await db.products.find_one({"barcode": barcode}, {"_id": 0})
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    return product

@api_router.get("/products/low-stock", response_model=List[Product])
async def get_low_stock_products(current_user: dict = Depends(get_current_user)):
    products = await db.products.find(
        {"$expr": {"$lte": ["$stock", "$min_stock_alert"]}},
        {"_id": 0}
    ).to_list(1000)
    return products

@api_router.get("/products/{product_id}", response_model=Product)
async def get_product(product_id: str, current_user: dict = Depends(get_current_user)):
    product = await db.products.find_one({"id": product_id}, {"_id": 0})
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    return product

@api_router.post("/products", response_model=Product)
async def create_product(data: ProductCreate, admin: dict = Depends(require_admin)):
    now = datetime.now(timezone.utc)
    items_per_box = data.items_per_box or 1
    warehouse_id = data.warehouse_id or "gudang_utama"
    
    # Stock input is in dus, convert to total satuan for product.stock
    total_satuan = data.stock * items_per_box
    
    product_data = data.model_dump()
    product_data.pop("warehouse_id", None)
    product_data["stock"] = total_satuan
    
    product_doc = {
        "id": str(uuid.uuid4()),
        **product_data,
        "box_price": data.box_price or 0,
        "items_per_box": items_per_box,
        "created_at": now.isoformat(),
        "updated_at": now.isoformat()
    }
    await db.products.insert_one(product_doc)
    
    # Create warehouse stock entries for all warehouses
    await ensure_warehouse_stock(product_doc["id"])
    
    # Put stock in the selected warehouse
    if data.stock > 0:
        wh = await db.warehouses.find_one({"id": warehouse_id}, {"_id": 0})
        if wh and wh.get("unit") == "satuan":
            await db.warehouse_stock.update_one(
                {"product_id": product_doc["id"], "warehouse_id": warehouse_id},
                {"$set": {"stock": total_satuan}}
            )
        else:
            await db.warehouse_stock.update_one(
                {"product_id": product_doc["id"], "warehouse_id": warehouse_id},
                {"$set": {"stock": data.stock}}
            )
        
        # Auto-create purchase record for the supplier
        if data.supplier_id:
            supplier = await db.suppliers.find_one({"id": data.supplier_id}, {"_id": 0})
            supplier_name = supplier["name"] if supplier else ""
            purchase_total = data.cost_price * data.stock  # cost_price per dus * jumlah dus
            paid = min(data.initial_paid or 0, purchase_total)
            remaining = purchase_total - paid
            
            purchase_doc = {
                "id": str(uuid.uuid4()),
                "supplier_id": data.supplier_id,
                "supplier_name": supplier_name,
                "items": [{
                    "product_id": product_doc["id"],
                    "product_name": data.name,
                    "quantity": data.stock,
                    "cost_price": data.cost_price,
                    "subtotal": purchase_total
                }],
                "total": purchase_total,
                "paid": paid,
                "remaining": remaining,
                "warehouse_id": warehouse_id,
                "notes": f"Pembelian awal - {data.name}" + (" (Tempo)" if remaining > 0 else ""),
                "created_at": now.isoformat(),
                "date": now.strftime("%Y-%m-%d")
            }
            await db.purchases.insert_one(purchase_doc)
    
    return product_doc

@api_router.put("/products/{product_id}", response_model=Product)
async def update_product(product_id: str, data: ProductCreate, admin: dict = Depends(require_admin)):
    update_data = data.model_dump()
    update_data.pop("warehouse_id", None)
    update_data.pop("stock", None)
    update_data.pop("initial_paid", None)
    update_data["updated_at"] = datetime.now(timezone.utc).isoformat()
    
    result = await db.products.find_one_and_update(
        {"id": product_id},
        {"$set": update_data},
        return_document=True,
        projection={"_id": 0}
    )
    if not result:
        raise HTTPException(status_code=404, detail="Product not found")
    
    await ensure_warehouse_stock(product_id)
    await recalculate_product_stock(product_id)
    await log_audit(admin, "update", "product", product_id, {"name": data.name})
    
    result = await db.products.find_one({"id": product_id}, {"_id": 0})
    return result

@api_router.delete("/products/{product_id}")
async def delete_product(product_id: str, admin: dict = Depends(require_admin)):
    product = await db.products.find_one({"id": product_id}, {"_id": 0})
    result = await db.products.delete_one({"id": product_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Product not found")
    await log_audit(admin, "delete", "product", product_id, {"name": product.get("name", "") if product else ""})
    return {"message": "Product deleted"}

@api_router.get("/pos/daily-summary")
async def get_pos_daily_summary(current_user: dict = Depends(get_current_user)):
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    transactions = await db.transactions.find({"date": today}, {"_id": 0}).to_list(10000)
    total_sales = sum(t.get("total", 0) for t in transactions)
    total_cost = 0
    for t in transactions:
        for item in t.get("items", []):
            total_cost += item.get("cost_price", 0) * item.get("quantity", 0)
    return {
        "count": len(transactions),
        "total_sales": total_sales,
        "total_cost": total_cost,
        "total_margin": total_sales - total_cost
    }

@api_router.post("/transactions", response_model=Transaction)
async def create_transaction(data: TransactionCreate, current_user: dict = Depends(get_current_user)):
    now = datetime.now(timezone.utc)
    transaction_number = f"TRX{now.strftime('%Y%m%d%H%M%S')}"
    
    for item in data.items:
        product = await db.products.find_one({"id": item.product_id})
        if not product:
            raise HTTPException(status_code=404, detail=f"Product {item.product_name} not found")
        
        # Deduct from Gudang Toko (satuan)
        toko_stock = await db.warehouse_stock.find_one({"product_id": item.product_id, "warehouse_id": "gudang_toko"})
        toko_qty = toko_stock["stock"] if toko_stock else 0
        if toko_qty < item.quantity:
            raise HTTPException(status_code=400, detail=f"Stok Gudang Toko tidak cukup untuk {item.product_name} (tersedia: {toko_qty} satuan)")
        
        await db.warehouse_stock.update_one(
            {"product_id": item.product_id, "warehouse_id": "gudang_toko"},
            {"$inc": {"stock": -item.quantity}}
        )
        await recalculate_product_stock(item.product_id)
    
    transaction_doc = {
        "id": str(uuid.uuid4()),
        "transaction_number": transaction_number,
        "cashier_id": current_user["id"],
        "cashier_name": current_user["name"],
        "items": [item.model_dump() for item in data.items],
        "subtotal": data.subtotal,
        "tax": data.tax,
        "discount": data.discount,
        "total": data.total,
        "payment_method": data.payment_method,
        "payment_amount": data.payment_amount,
        "change": data.change,
        "created_at": now.isoformat(),
        "date": now.strftime("%Y-%m-%d")
    }
    
    await db.transactions.insert_one(transaction_doc)
    return transaction_doc

@api_router.get("/transactions", response_model=List[Transaction])
async def get_transactions(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    query = {}
    if start_date and end_date:
        query["date"] = {"$gte": start_date, "$lte": end_date}
    
    transactions = await db.transactions.find(query, {"_id": 0}).sort("created_at", -1).to_list(1000)
    return transactions

@api_router.get("/transactions/{transaction_id}", response_model=Transaction)
async def get_transaction(transaction_id: str, current_user: dict = Depends(get_current_user)):
    transaction = await db.transactions.find_one({"id": transaction_id}, {"_id": 0})
    if not transaction:
        raise HTTPException(status_code=404, detail="Transaction not found")
    return transaction

@api_router.get("/reports/daily")
async def get_daily_report(date: str, admin: dict = Depends(require_admin)):
    transactions = await db.transactions.find({"date": date}, {"_id": 0}).to_list(1000)
    
    total_sales = sum(t["total"] for t in transactions)
    total_transactions = len(transactions)
    
    payment_methods = {}
    for t in transactions:
        method = t["payment_method"]
        payment_methods[method] = payment_methods.get(method, 0) + t["total"]
    
    return {
        "date": date,
        "total_sales": total_sales,
        "total_transactions": total_transactions,
        "payment_methods": payment_methods,
        "transactions": transactions
    }

@api_router.get("/reports/monthly")
async def get_monthly_report(year: int, month: int, admin: dict = Depends(require_admin)):
    start_date = f"{year}-{month:02d}-01"
    if month == 12:
        end_date = f"{year+1}-01-01"
    else:
        end_date = f"{year}-{month+1:02d}-01"
    
    transactions = await db.transactions.find(
        {"date": {"$gte": start_date, "$lt": end_date}},
        {"_id": 0}
    ).to_list(10000)
    
    total_sales = sum(t["total"] for t in transactions)
    total_transactions = len(transactions)
    
    daily_sales = {}
    for t in transactions:
        date = t["date"]
        daily_sales[date] = daily_sales.get(date, 0) + t["total"]
    
    return {
        "year": year,
        "month": month,
        "total_sales": total_sales,
        "total_transactions": total_transactions,
        "daily_sales": daily_sales
    }


@api_router.get("/reports/sales-detail")
async def get_sales_detail_report(date_from: str, date_to: str, admin: dict = Depends(require_admin)):
    """Detailed per-product sales report from POS transactions"""
    transactions = await db.transactions.find(
        {"date": {"$gte": date_from, "$lte": date_to}}, {"_id": 0}
    ).to_list(10000)

    # Build product lookup for cost data
    product_ids = set()
    for t in transactions:
        for item in t.get("items", []):
            product_ids.add(item["product_id"])
    products = {}
    if product_ids:
        for p in await db.products.find({"id": {"$in": list(product_ids)}}, {"_id": 0}).to_list(10000):
            products[p["id"]] = p

    rows = []
    total_sales = 0
    total_cost = 0
    total_margin = 0
    for t in transactions:
        for item in t.get("items", []):
            product = products.get(item["product_id"], {})
            ipb = product.get("items_per_box", 1) or 1
            cost_per_unit = (product.get("cost_price", 0) or 0) / ipb
            sell_price = item["price"]
            qty = item["quantity"]
            total_sell = sell_price * qty
            total_cost_item = cost_per_unit * qty
            margin = total_sell - total_cost_item
            rows.append({
                "date": t["date"],
                "time": t.get("created_at", ""),
                "transaction_number": t.get("transaction_number", ""),
                "product_name": item["product_name"],
                "quantity": qty,
                "cost_price": round(cost_per_unit, 2),
                "sell_price": sell_price,
                "total_sell": total_sell,
                "total_cost": round(total_cost_item, 2),
                "margin": round(margin, 2)
            })
            total_sales += total_sell
            total_cost += total_cost_item
            total_margin += margin

    return {
        "date_from": date_from,
        "date_to": date_to,
        "rows": rows,
        "total_sales": round(total_sales, 2),
        "total_cost": round(total_cost, 2),
        "total_margin": round(total_margin, 2),
        "total_items": len(rows)
    }

@api_router.get("/reports/partner-sales-detail")
async def get_partner_sales_detail_report(date_from: str, date_to: str, admin: dict = Depends(require_admin)):
    """Detailed per-product partner order sales report"""
    orders = await db.partner_orders.find(
        {"date": {"$gte": date_from, "$lte": date_to}}, {"_id": 0}
    ).to_list(10000)

    product_ids = set()
    for o in orders:
        for item in o.get("items", []):
            product_ids.add(item["product_id"])
    products = {}
    if product_ids:
        for p in await db.products.find({"id": {"$in": list(product_ids)}}, {"_id": 0}).to_list(10000):
            products[p["id"]] = p

    warehouses_list = await db.warehouses.find({}, {"_id": 0}).to_list(100)
    wh_map = {w["id"]: w for w in warehouses_list}

    rows = []
    total_sales = 0
    total_cost = 0
    total_margin = 0
    for o in orders:
        for item in o.get("items", []):
            product = products.get(item["product_id"], {})
            ipb = product.get("items_per_box", 1) or 1
            item_wh_id = item.get("warehouse_id") or o.get("warehouse_id", "gudang_utama")
            wh = wh_map.get(item_wh_id, {})
            wh_unit = wh.get("unit", "satuan")

            # Convert quantity to satuan for consistent reporting
            qty_satuan = item["quantity"] * ipb if wh_unit == "dus" else item["quantity"]

            cost_per_unit = (product.get("cost_price", 0) or 0) / ipb
            sell_price_per_unit = item["price"] / ipb if wh_unit == "dus" else item["price"]
            total_sell = item["price"] * item["quantity"]
            total_cost_item = cost_per_unit * qty_satuan
            margin = total_sell - total_cost_item

            rows.append({
                "date": o["date"],
                "order_number": o.get("order_number", ""),
                "partner_store_name": o.get("partner_store_name", ""),
                "product_name": item["product_name"],
                "quantity_original": item["quantity"],
                "unit_original": wh_unit,
                "quantity_satuan": qty_satuan,
                "cost_price_per_unit": round(cost_per_unit, 2),
                "sell_price_per_unit": round(sell_price_per_unit, 2),
                "total_sell": round(total_sell, 2),
                "total_cost": round(total_cost_item, 2),
                "margin": round(margin, 2)
            })
            total_sales += total_sell
            total_cost += total_cost_item
            total_margin += margin

    return {
        "date_from": date_from,
        "date_to": date_to,
        "rows": rows,
        "total_sales": round(total_sales, 2),
        "total_cost": round(total_cost, 2),
        "total_margin": round(total_margin, 2),
        "total_items": len(rows)
    }

@api_router.get("/reports/best-sellers")
async def get_best_sellers_report(date_from: str, date_to: str, admin: dict = Depends(require_admin)):
    """Product ranking by total quantity sold and revenue (POS + Partner combined)"""
    # Fetch POS transactions
    transactions = await db.transactions.find(
        {"date": {"$gte": date_from, "$lte": date_to}}, {"_id": 0}
    ).to_list(10000)

    # Fetch partner orders
    orders = await db.partner_orders.find(
        {"date": {"$gte": date_from, "$lte": date_to}}, {"_id": 0}
    ).to_list(10000)

    # Build product lookup
    all_product_ids = set()
    for t in transactions:
        for item in t.get("items", []):
            all_product_ids.add(item["product_id"])
    for o in orders:
        for item in o.get("items", []):
            all_product_ids.add(item["product_id"])

    products = {}
    if all_product_ids:
        for p in await db.products.find({"id": {"$in": list(all_product_ids)}}, {"_id": 0}).to_list(10000):
            products[p["id"]] = p

    warehouses_list = await db.warehouses.find({}, {"_id": 0}).to_list(100)
    wh_map = {w["id"]: w for w in warehouses_list}

    # Aggregate per product
    agg = {}
    for t in transactions:
        for item in t.get("items", []):
            pid = item["product_id"]
            if pid not in agg:
                agg[pid] = {"product_name": item["product_name"], "qty_pos": 0, "qty_partner": 0, "revenue_pos": 0, "revenue_partner": 0, "cost_total": 0}
            agg[pid]["qty_pos"] += item["quantity"]
            agg[pid]["revenue_pos"] += item["subtotal"]
            product = products.get(pid, {})
            ipb = product.get("items_per_box", 1) or 1
            cost_per_unit = (product.get("cost_price", 0) or 0) / ipb
            agg[pid]["cost_total"] += cost_per_unit * item["quantity"]

    for o in orders:
        for item in o.get("items", []):
            pid = item["product_id"]
            if pid not in agg:
                agg[pid] = {"product_name": item["product_name"], "qty_pos": 0, "qty_partner": 0, "revenue_pos": 0, "revenue_partner": 0, "cost_total": 0}
            product = products.get(pid, {})
            ipb = product.get("items_per_box", 1) or 1
            item_wh_id = item.get("warehouse_id") or o.get("warehouse_id", "gudang_utama")
            wh = wh_map.get(item_wh_id, {})
            wh_unit = wh.get("unit", "satuan")
            qty_satuan = item["quantity"] * ipb if wh_unit == "dus" else item["quantity"]
            agg[pid]["qty_partner"] += qty_satuan
            agg[pid]["revenue_partner"] += item["subtotal"]
            cost_per_unit = (product.get("cost_price", 0) or 0) / ipb
            agg[pid]["cost_total"] += cost_per_unit * qty_satuan

    # Build ranked list
    rows = []
    for pid, data in agg.items():
        total_qty = data["qty_pos"] + data["qty_partner"]
        total_revenue = data["revenue_pos"] + data["revenue_partner"]
        total_cost = data["cost_total"]
        margin = total_revenue - total_cost
        rows.append({
            "product_id": pid,
            "product_name": data["product_name"],
            "qty_pos": data["qty_pos"],
            "qty_partner": data["qty_partner"],
            "qty_total": total_qty,
            "revenue_pos": round(data["revenue_pos"], 2),
            "revenue_partner": round(data["revenue_partner"], 2),
            "revenue_total": round(total_revenue, 2),
            "cost_total": round(total_cost, 2),
            "margin": round(margin, 2)
        })

    # Sort by revenue (highest first)
    rows.sort(key=lambda x: x["revenue_total"], reverse=True)

    return {
        "date_from": date_from,
        "date_to": date_to,
        "rows": rows,
        "total_products": len(rows)
    }


@api_router.get("/dashboard/stats")
async def get_dashboard_stats(admin: dict = Depends(require_admin)):
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    
    today_transactions = await db.transactions.find({"date": today}, {"_id": 0}).to_list(1000)
    today_sales = sum(t["total"] for t in today_transactions)
    
    total_products = await db.products.count_documents({})
    low_stock_count = await db.products.count_documents(
        {"$expr": {"$lte": ["$stock", "$min_stock_alert"]}}
    )
    
    last_7_days = []
    for i in range(6, -1, -1):
        date = (datetime.now(timezone.utc) - timedelta(days=i)).strftime("%Y-%m-%d")
        trans = await db.transactions.find({"date": date}, {"_id": 0}).to_list(1000)
        daily_total = sum(t["total"] for t in trans)
        last_7_days.append({"date": date, "sales": daily_total})
    
    # Inventory value per warehouse
    products = await db.products.find({}, {"_id": 0}).to_list(100000)
    products_map = {p["id"]: p for p in products}
    warehouses = await db.warehouses.find({}, {"_id": 0}).to_list(100)
    warehouse_stocks = await db.warehouse_stock.find({}, {"_id": 0}).to_list(100000)
    
    warehouse_values = {}
    for ws in warehouse_stocks:
        pid = ws.get("product_id")
        wid = ws.get("warehouse_id")
        stock = ws.get("stock", 0)
        product = products_map.get(pid)
        if product and stock > 0:
            hpp_per_unit = product.get("cost_price", 0) / max(1, product.get("items_per_box", 1))
            wh = next((w for w in warehouses if w.get("id") == wid), None)
            if wh and wh.get("unit") == "dus":
                value = stock * product.get("cost_price", 0)
            else:
                value = stock * hpp_per_unit
            warehouse_values[wid] = warehouse_values.get(wid, 0) + value
    
    warehouse_chart = []
    total_inventory = 0
    for w in warehouses:
        val = warehouse_values.get(w["id"], 0)
        total_inventory += val
        if val > 0:
            warehouse_chart.append({"name": w.get("name", w["id"]), "value": round(val)})
    
    # Stock movement today (purchases = in, transactions = out)
    today_purchases = await db.purchases.find({"date": today}, {"_id": 0}).to_list(1000)
    stock_in_today = sum(
        sum(item.get("subtotal", 0) for item in p.get("items", []))
        for p in today_purchases
    )
    stock_out_today = sum(
        sum(item.get("cost_price", 0) * item.get("quantity", 0) for item in t.get("items", []))
        for t in today_transactions
    )
    
    # Supplier debt
    unpaid_purchases = await db.purchases.find({"remaining": {"$gt": 0}}, {"_id": 0}).to_list(10000)
    total_supplier_debt = sum(p.get("remaining", 0) for p in unpaid_purchases)
    
    # Partner debt
    partner_stores = await db.partner_stores.find({}, {"_id": 0, "current_debt": 1}).to_list(10000)
    total_partner_debt = sum(ps.get("current_debt", 0) for ps in partner_stores)
    
    # Total assets = inventory + partner debt - supplier debt
    total_assets = total_inventory + total_partner_debt - total_supplier_debt
    
    # Due this month (per invoice)
    now = datetime.now(timezone.utc)
    month_start = now.strftime("%Y-%m-01")
    month_end = f"{now.strftime('%Y-%m')}-31"
    due_this_month = await db.purchases.find(
        {"due_date": {"$gte": month_start, "$lte": month_end}, "remaining": {"$gt": 0}},
        {"_id": 0}
    ).to_list(100)
    
    return {
        "today_sales": today_sales,
        "today_transactions": len(today_transactions),
        "total_products": total_products,
        "low_stock_count": low_stock_count,
        "last_7_days": last_7_days,
        "inventory_value": total_inventory,
        "warehouse_chart": warehouse_chart,
        "stock_in_today": stock_in_today,
        "stock_out_today": stock_out_today,
        "total_supplier_debt": total_supplier_debt,
        "total_partner_debt": total_partner_debt,
        "total_assets": total_assets,
        "due_this_month": [{k: v for k, v in d.items() if k != "_id"} for d in due_this_month]
    }

@api_router.get("/dashboard/due-calendar")
async def get_due_calendar(year: int, month: int, admin: dict = Depends(require_admin)):
    month_start = f"{year:04d}-{month:02d}-01"
    month_end = f"{year:04d}-{month:02d}-31"
    purchases = await db.purchases.find(
        {"due_date": {"$gte": month_start, "$lte": month_end}, "remaining": {"$gt": 0}},
        {"_id": 0}
    ).to_list(1000)
    return [{k: v for k, v in p.items() if k != "_id"} for p in purchases]

@api_router.get("/users", response_model=List[User])
async def get_users(admin: dict = Depends(require_admin)):
    users = await db.users.find({}, {"_id": 0, "password": 0}).sort("name", 1).to_list(1000)
    return users

@api_router.delete("/users/{user_id}")
async def delete_user(user_id: str, admin: dict = Depends(require_admin)):
    result = await db.users.delete_one({"id": user_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="User not found")
    return {"message": "User deleted"}

# ============= STORE SETTINGS =============
@api_router.get("/store-settings")
async def get_store_settings(current_user: dict = Depends(get_current_user)):
    settings = await db.store_settings.find_one({}, {"_id": 0})
    if not settings:
        # Return default settings
        return {
            "store_name": "UD. AQIN TANI",
            "address": "Jl. Contoh No. 123, Kota Anda",
            "phone": "(021) 1234-5678",
            "email": "info@aqintani.com",
            "logo_url": ""
        }
    return settings

@api_router.put("/store-settings")
async def update_store_settings(settings: StoreSettings, admin: dict = Depends(require_admin)):
    settings_doc = settings.model_dump()
    await db.store_settings.delete_many({})
    await db.store_settings.insert_one(settings_doc)
    return settings_doc

# ============= PURCHASES (from Suppliers) =============
@api_router.get("/purchases")
async def get_purchases(
    supplier_id: Optional[str] = None,
    admin: dict = Depends(require_admin)
):
    query = {}
    if supplier_id:
        query["supplier_id"] = supplier_id
    purchases = await db.purchases.find(query, {"_id": 0}).sort("created_at", -1).to_list(1000)
    return purchases

@api_router.put("/purchases/{purchase_id}/due-date")
async def set_purchase_due_date(purchase_id: str, due_date: str = "", admin: dict = Depends(require_admin)):
    update = {"$set": {"due_date": due_date}} if due_date else {"$unset": {"due_date": ""}}
    result = await db.purchases.update_one({"id": purchase_id}, update)
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Pembelian tidak ditemukan")
    return {"message": "Tanggal jatuh tempo berhasil diperbarui"}

@api_router.post("/purchases/{purchase_id}/return")
async def process_purchase_return(purchase_id: str, request: Request, admin: dict = Depends(require_admin)):
    body = await request.json()
    items = body.get("items", [])
    notes = body.get("notes", "")
    warehouse_id = body.get("warehouse_id", "gudang_utama")

    purchase = await db.purchases.find_one({"id": purchase_id}, {"_id": 0})
    if not purchase:
        raise HTTPException(status_code=404, detail="Pembelian tidak ditemukan")

    now = datetime.now(timezone.utc)
    total_return_value = 0
    return_items = []

    for item in items:
        product_id = item.get("product_id")
        qty = int(item.get("quantity", 0))
        if qty <= 0:
            continue
        product = await db.products.find_one({"id": product_id}, {"_id": 0})
        if not product:
            continue
        purchase_item = next((pi for pi in purchase.get("items", []) if pi["product_id"] == product_id), None)
        cost = purchase_item["cost_price"] if purchase_item else product.get("cost_price", 0)
        subtotal = qty * cost
        total_return_value += subtotal
        return_items.append({
            "product_id": product_id,
            "product_name": product["name"],
            "quantity": qty,
            "cost_price": cost,
            "subtotal": subtotal
        })
        # Reduce warehouse stock
        wh_stock = await db.warehouse_stock.find_one({"product_id": product_id, "warehouse_id": warehouse_id})
        if wh_stock:
            new_stock = max(0, wh_stock.get("stock", 0) - qty)
            await db.warehouse_stock.update_one(
                {"product_id": product_id, "warehouse_id": warehouse_id},
                {"$set": {"stock": new_stock}}
            )
        await recalculate_product_stock(product_id)

    if not return_items:
        raise HTTPException(status_code=400, detail="Tidak ada item yang diretur")

    # Reduce purchase remaining
    new_remaining = max(0, purchase.get("remaining", 0) - total_return_value)
    await db.purchases.update_one(
        {"id": purchase_id},
        {"$set": {"remaining": new_remaining}}
    )

    # Create return record
    return_doc = {
        "id": str(uuid.uuid4()),
        "type": "supplier_return",
        "purchase_id": purchase_id,
        "supplier_id": purchase.get("supplier_id"),
        "supplier_name": purchase.get("supplier_name"),
        "items": return_items,
        "total": total_return_value,
        "warehouse_id": warehouse_id,
        "notes": notes or f"Retur ke {purchase.get('supplier_name', '')}",
        "created_at": now.isoformat(),
        "date": now.strftime("%Y-%m-%d")
    }
    await db.payment_history.insert_one(return_doc)

    return {
        "message": f"Retur berhasil. Hutang berkurang {total_return_value}",
        "return_value": total_return_value,
        "new_remaining": new_remaining
    }

@api_router.post("/purchases", response_model=Purchase)
async def create_purchase(data: PurchaseCreate, admin: dict = Depends(require_admin)):
    supplier = await db.suppliers.find_one({"id": data.supplier_id}, {"_id": 0})
    if not supplier:
        raise HTTPException(status_code=404, detail="Supplier not found")
    
    warehouse_id = data.warehouse_id or "gudang_utama"
    
    now = datetime.now(timezone.utc)
    purchase_doc = {
        "id": str(uuid.uuid4()),
        "supplier_id": data.supplier_id,
        "supplier_name": supplier["name"],
        "items": data.items,
        "total": data.total,
        "paid": data.paid,
        "remaining": data.total - data.paid,
        "warehouse_id": warehouse_id,
        "notes": data.notes,
        "created_at": now.isoformat(),
        "date": now.strftime("%Y-%m-%d")
    }
    
    await db.purchases.insert_one(purchase_doc)
    
    # Add purchased quantities to warehouse stock
    for item in data.items:
        pid = item.get("product_id")
        qty = item.get("quantity", 0)
        if pid and qty > 0:
            await ensure_warehouse_stock(pid)
            await db.warehouse_stock.update_one(
                {"product_id": pid, "warehouse_id": warehouse_id},
                {"$inc": {"stock": qty}},
                upsert=True
            )
            await recalculate_product_stock(pid)
    
    return purchase_doc

@api_router.post("/purchases/{purchase_id}/pay")
async def pay_purchase(purchase_id: str, amount: float, admin: dict = Depends(require_admin)):
    purchase = await db.purchases.find_one({"id": purchase_id})
    if not purchase:
        raise HTTPException(status_code=404, detail="Purchase not found")
    
    new_paid = purchase["paid"] + amount
    new_remaining = purchase["total"] - new_paid
    
    result = await db.purchases.find_one_and_update(
        {"id": purchase_id},
        {"$set": {"paid": new_paid, "remaining": new_remaining}},
        return_document=True,
        projection={"_id": 0}
    )
    
    # Record payment history
    await db.payment_history.insert_one({
        "id": str(uuid.uuid4()),
        "type": "supplier",
        "reference_id": purchase_id,
        "supplier_id": purchase.get("supplier_id", ""),
        "supplier_name": purchase.get("supplier_name", ""),
        "amount": amount,
        "notes": f"Pembayaran pembelian {purchase.get('id', '')}",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "date": datetime.now(timezone.utc).strftime("%Y-%m-%d")
    })
    
    return result

@api_router.delete("/purchases/{purchase_id}")
async def delete_purchase(purchase_id: str, admin: dict = Depends(require_admin)):
    result = await db.purchases.delete_one({"id": purchase_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Purchase not found")
    return {"message": "Purchase deleted"}

# ============= PAYMENT HISTORY =============
@api_router.get("/payment-history")
async def get_payment_history(
    type: Optional[str] = None,
    reference_id: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    query = {}
    if type:
        query["type"] = type
    if reference_id:
        query["reference_id"] = reference_id
    history = await db.payment_history.find(query, {"_id": 0}).sort("created_at", -1).to_list(1000)
    return history

@api_router.post("/partner-stores/{store_id}/pay")
async def pay_partner_store(store_id: str, amount: float, notes: str = "", admin: dict = Depends(require_admin)):
    store = await db.partner_stores.find_one({"id": store_id})
    if not store:
        raise HTTPException(status_code=404, detail="Partner store not found")
    
    new_debt = max(0, store.get("current_debt", 0) - amount)
    await db.partner_stores.update_one(
        {"id": store_id},
        {"$set": {"current_debt": new_debt}}
    )
    
    payment_doc = {
        "id": str(uuid.uuid4()),
        "type": "partner_store",
        "reference_id": store_id,
        "partner_store_name": store["name"],
        "amount": amount,
        "notes": notes or f"Setoran dari {store['name']}",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "date": datetime.now(timezone.utc).strftime("%Y-%m-%d")
    }
    await db.payment_history.insert_one(payment_doc)
    
    return {"message": "Payment recorded", "new_debt": new_debt, "payment": {k: v for k, v in payment_doc.items() if k != "_id"}}

@api_router.get("/partner-stores/{store_id}/payments")
async def get_partner_store_payments(store_id: str, current_user: dict = Depends(get_current_user)):
    payments = await db.payment_history.find(
        {"type": "partner_store", "reference_id": store_id},
        {"_id": 0}
    ).sort("created_at", -1).to_list(1000)
    return payments

@api_router.get("/suppliers/{supplier_id}/payments")
async def get_supplier_payments(supplier_id: str, current_user: dict = Depends(get_current_user)):
    payments = await db.payment_history.find(
        {"type": "supplier", "supplier_id": supplier_id},
        {"_id": 0}
    ).sort("created_at", -1).to_list(1000)
    return payments

# ============= SUPPLIER DEBT SUMMARY =============
@api_router.get("/suppliers/{supplier_id}/debt")
async def get_supplier_debt(supplier_id: str, admin: dict = Depends(require_admin)):
    purchases = await db.purchases.find({"supplier_id": supplier_id}, {"_id": 0}).to_list(1000)
    total_purchases = sum(p.get("total", 0) for p in purchases)
    total_paid = sum(p.get("paid", 0) for p in purchases)
    total_remaining = total_purchases - total_paid
    return {
        "supplier_id": supplier_id,
        "total_purchases": total_purchases,
        "total_paid": total_paid,
        "total_remaining": total_remaining,
        "purchase_count": len(purchases)
    }

# ============= SUPPLIER DEBT MANAGEMENT =============
@api_router.get("/supplier-debts/overview")
async def get_supplier_debts_overview(admin: dict = Depends(require_admin)):
    """Returns all suppliers with their aggregated debt summary."""
    suppliers = await db.suppliers.find({}, {"_id": 0}).sort("name", 1).to_list(1000)
    purchases = await db.purchases.find({}, {"_id": 0}).to_list(10000)

    # Group purchases by supplier
    supplier_purchases = {}
    for p in purchases:
        sid = p.get("supplier_id", "")
        if sid not in supplier_purchases:
            supplier_purchases[sid] = []
        supplier_purchases[sid].append(p)

    result = []
    for s in suppliers:
        sid = s["id"]
        sp = supplier_purchases.get(sid, [])
        total_purchases = sum(p.get("total", 0) for p in sp)
        total_paid = sum(p.get("paid", 0) for p in sp)
        outstanding = total_purchases - total_paid
        unpaid_purchases = [p for p in sp if p.get("remaining", 0) > 0]
        last_purchase_date = max((p.get("date", "") for p in sp), default="") if sp else ""

        # Calculate oldest unpaid purchase date as "due indicator"
        oldest_unpaid_date = ""
        if unpaid_purchases:
            oldest_unpaid_date = min(p.get("date", "") for p in unpaid_purchases)

        result.append({
            "supplier_id": sid,
            "supplier_name": s["name"],
            "phone": s.get("phone", ""),
            "contact_person": s.get("contact_person", ""),
            "total_purchases": total_purchases,
            "total_paid": total_paid,
            "outstanding": outstanding,
            "purchase_count": len(sp),
            "unpaid_count": len(unpaid_purchases),
            "last_purchase_date": last_purchase_date,
            "oldest_unpaid_date": oldest_unpaid_date,
        })

    # Sort: suppliers with debt first, then by outstanding descending
    result.sort(key=lambda x: (-1 if x["outstanding"] > 0 else 0, -x["outstanding"]))
    return result

@api_router.get("/supplier-debts/summary")
async def get_supplier_debts_summary(admin: dict = Depends(require_admin)):
    """Aggregate summary: total debt, suppliers with debt, payments this month."""
    purchases = await db.purchases.find({}, {"_id": 0}).to_list(10000)

    total_debt = sum(max(0, p.get("remaining", 0)) for p in purchases)
    suppliers_with_debt = len(set(
        p["supplier_id"] for p in purchases if p.get("remaining", 0) > 0
    ))

    # Payments this month
    now = datetime.now(timezone.utc)
    month_start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0).isoformat()
    payments_this_month = await db.payment_history.find(
        {"type": "supplier", "created_at": {"$gte": month_start}},
        {"_id": 0}
    ).to_list(10000)
    total_paid_this_month = sum(p.get("amount", 0) for p in payments_this_month)

    # Overdue count (purchases older than 30 days with remaining > 0)
    thirty_days_ago = (now - timedelta(days=30)).strftime("%Y-%m-%d")
    overdue_purchases = [
        p for p in purchases
        if p.get("remaining", 0) > 0 and p.get("date", "") <= thirty_days_ago
    ]
    overdue_total = sum(p.get("remaining", 0) for p in overdue_purchases)

    return {
        "total_debt": total_debt,
        "suppliers_with_debt": suppliers_with_debt,
        "total_paid_this_month": total_paid_this_month,
        "payments_this_month_count": len(payments_this_month),
        "overdue_count": len(overdue_purchases),
        "overdue_total": overdue_total,
    }

@api_router.get("/supplier-debts/{supplier_id}/purchases")
async def get_supplier_unpaid_purchases(supplier_id: str, show_all: bool = False, admin: dict = Depends(require_admin)):
    """Returns purchases for a supplier, optionally filtering to unpaid only."""
    query = {"supplier_id": supplier_id}
    if not show_all:
        query["remaining"] = {"$gt": 0}
    purchases = await db.purchases.find(query, {"_id": 0}).sort("date", -1).to_list(1000)
    return purchases

@api_router.post("/supplier-debts/{supplier_id}/pay")
async def pay_supplier_debt(supplier_id: str, amount: float, notes: str = "", admin: dict = Depends(require_admin)):
    """Pay supplier debt — distributes payment to oldest unpaid purchases first."""
    supplier = await db.suppliers.find_one({"id": supplier_id}, {"_id": 0})
    if not supplier:
        raise HTTPException(status_code=404, detail="Supplier not found")

    # Get unpaid purchases sorted by date (oldest first)
    unpaid = await db.purchases.find(
        {"supplier_id": supplier_id, "remaining": {"$gt": 0}},
        {"_id": 0}
    ).sort("date", 1).to_list(1000)

    if not unpaid:
        raise HTTPException(status_code=400, detail="Tidak ada hutang untuk supplier ini")

    remaining_amount = amount
    paid_purchases = []

    for purchase in unpaid:
        if remaining_amount <= 0:
            break
        purchase_remaining = purchase.get("remaining", 0)
        pay_this = min(remaining_amount, purchase_remaining)

        new_paid = purchase["paid"] + pay_this
        new_remaining = purchase["total"] - new_paid

        await db.purchases.update_one(
            {"id": purchase["id"]},
            {"$set": {"paid": new_paid, "remaining": max(0, new_remaining)}}
        )
        paid_purchases.append({"purchase_id": purchase["id"], "amount": pay_this})
        remaining_amount -= pay_this

    # Record payment history
    payment_doc = {
        "id": str(uuid.uuid4()),
        "type": "supplier",
        "reference_id": supplier_id,
        "supplier_id": supplier_id,
        "supplier_name": supplier["name"],
        "amount": amount,
        "notes": notes or f"Pembayaran hutang ke {supplier['name']}",
        "paid_purchases": paid_purchases,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "date": datetime.now(timezone.utc).strftime("%Y-%m-%d")
    }
    await db.payment_history.insert_one(payment_doc)

    # Recalculate remaining debt
    all_purchases = await db.purchases.find({"supplier_id": supplier_id}, {"_id": 0}).to_list(1000)
    new_total_debt = sum(max(0, p.get("remaining", 0)) for p in all_purchases)

    return {
        "message": "Pembayaran berhasil dicatat",
        "amount_paid": amount,
        "new_total_debt": new_total_debt,
        "paid_purchases": paid_purchases,
        "payment_id": payment_doc["id"]
    }

@api_router.get("/supplier-debts/export-csv")
async def export_supplier_debts_csv(admin: dict = Depends(require_admin)):
    """Export supplier debts to CSV."""
    from io import StringIO
    import csv
    from starlette.responses import StreamingResponse

    suppliers = await db.suppliers.find({}, {"_id": 0}).sort("name", 1).to_list(1000)
    purchases = await db.purchases.find({}, {"_id": 0}).to_list(10000)

    supplier_purchases = {}
    for p in purchases:
        sid = p.get("supplier_id", "")
        if sid not in supplier_purchases:
            supplier_purchases[sid] = []
        supplier_purchases[sid].append(p)

    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(["Supplier", "Kontak", "Telepon", "Jml Pembelian", "Total Pembelian", "Total Dibayar", "Sisa Hutang", "Pembelian Belum Lunas", "Tanggal Pembelian Terakhir"])

    for s in suppliers:
        sid = s["id"]
        sp = supplier_purchases.get(sid, [])
        total = sum(p.get("total", 0) for p in sp)
        paid = sum(p.get("paid", 0) for p in sp)
        outstanding = total - paid
        unpaid_count = len([p for p in sp if p.get("remaining", 0) > 0])
        last_date = max((p.get("date", "") for p in sp), default="-") if sp else "-"

        if outstanding > 0 or len(sp) > 0:
            writer.writerow([
                s["name"], s.get("contact_person", ""), s.get("phone", ""),
                len(sp), f"{total:.0f}", f"{paid:.0f}", f"{outstanding:.0f}",
                unpaid_count, last_date
            ])

    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=hutang_supplier_{datetime.now(timezone.utc).strftime('%Y%m%d')}.csv"}
    )

# ============= PARTNER STORES (Kios Mitra) =============
@api_router.get("/partner-stores", response_model=List[PartnerStore])
async def get_partner_stores(current_user: dict = Depends(get_current_user)):
    stores = await db.partner_stores.find({}, {"_id": 0}).sort("name", 1).to_list(1000)
    return stores

@api_router.post("/partner-stores", response_model=PartnerStore)
async def create_partner_store(data: PartnerStoreCreate, admin: dict = Depends(require_admin)):
    store_doc = {
        "id": str(uuid.uuid4()),
        **data.model_dump(),
        "current_debt": 0,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.partner_stores.insert_one(store_doc)
    return store_doc

@api_router.put("/partner-stores/{store_id}", response_model=PartnerStore)
async def update_partner_store(store_id: str, data: PartnerStoreCreate, admin: dict = Depends(require_admin)):
    result = await db.partner_stores.find_one_and_update(
        {"id": store_id},
        {"$set": data.model_dump()},
        return_document=True,
        projection={"_id": 0}
    )
    if not result:
        raise HTTPException(status_code=404, detail="Partner store not found")
    return result

@api_router.delete("/partner-stores/{store_id}")
async def delete_partner_store(store_id: str, admin: dict = Depends(require_admin)):
    result = await db.partner_stores.delete_one({"id": store_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Partner store not found")
    return {"message": "Partner store deleted"}

# ============= PARTNER ORDERS =============
@api_router.get("/partner-orders", response_model=List[PartnerOrder])
async def get_partner_orders(
    partner_store_id: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    query = {}
    if partner_store_id:
        query["partner_store_id"] = partner_store_id
    orders = await db.partner_orders.find(query, {"_id": 0}).sort("created_at", -1).to_list(1000)
    return orders

@api_router.get("/partner-orders/{order_id}")
async def get_partner_order(order_id: str, current_user: dict = Depends(get_current_user)):
    order = await db.partner_orders.find_one({"id": order_id}, {"_id": 0})
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    # Attach partner store current_debt for receipt
    partner_store = await db.partner_stores.find_one({"id": order["partner_store_id"]}, {"_id": 0})
    order["partner_debt_total"] = partner_store.get("current_debt", 0) if partner_store else 0
    order["partner_phone"] = partner_store.get("phone", "") if partner_store else ""
    order["partner_address"] = partner_store.get("address", "") if partner_store else ""
    return order


@api_router.post("/partner-orders", response_model=PartnerOrder)
async def create_partner_order(data: PartnerOrderCreate, current_user: dict = Depends(get_current_user)):
    partner_store = await db.partner_stores.find_one({"id": data.partner_store_id}, {"_id": 0})
    if not partner_store:
        raise HTTPException(status_code=404, detail="Partner store not found")
    
    canvasser_name = ""
    if data.canvasser_id:
        canvasser = await db.canvassers.find_one({"id": data.canvasser_id}, {"_id": 0})
        if canvasser:
            canvasser_name = canvasser["name"]
    
    # Deduct stock from each item's warehouse (or fallback to order-level warehouse)
    default_wh_id = data.warehouse_id or "gudang_utama"
    for item in data.items:
        item_wh_id = item.warehouse_id or default_wh_id
        wh = await db.warehouses.find_one({"id": item_wh_id}, {"_id": 0})
        if wh:
            wh_stock = await db.warehouse_stock.find_one({"product_id": item.product_id, "warehouse_id": item_wh_id})
            current_qty = wh_stock["stock"] if wh_stock else 0
            if current_qty < item.quantity:
                wh_name = wh.get("name", item_wh_id)
                raise HTTPException(status_code=400, detail=f"Stok {wh_name} tidak cukup untuk {item.product_name}")
            await db.warehouse_stock.update_one(
                {"product_id": item.product_id, "warehouse_id": item_wh_id},
                {"$inc": {"stock": -item.quantity}}
            )
            await recalculate_product_stock(item.product_id)
    
    now = datetime.now(timezone.utc)
    order_number = f"PO{now.strftime('%Y%m%d%H%M%S')}"
    
    order_doc = {
        "id": str(uuid.uuid4()),
        "order_number": order_number,
        "partner_store_id": data.partner_store_id,
        "partner_store_name": partner_store["name"],
        "items": [item.model_dump() for item in data.items],
        "subtotal": data.subtotal,
        "discount": data.discount,
        "total": data.total,
        "paid": data.paid,
        "remaining": data.total - data.paid,
        "delivery_date": data.delivery_date or now.strftime("%Y-%m-%d"),
        "canvasser_id": data.canvasser_id or "",
        "canvasser_name": canvasser_name,
        "warehouse_id": default_wh_id,
        "status": "pending",
        "created_at": now.isoformat(),
        "date": now.strftime("%Y-%m-%d")
    }
    
    await db.partner_orders.insert_one(order_doc)
    
    # Update partner store debt
    new_debt = partner_store.get("current_debt", 0) + order_doc["remaining"]
    await db.partner_stores.update_one(
        {"id": data.partner_store_id},
        {"$set": {"current_debt": new_debt}}
    )
    
    return order_doc

@api_router.put("/partner-orders/{order_id}/status")
async def update_order_status(order_id: str, order_status: str = Query(..., alias="status"), current_user: dict = Depends(get_current_user)):
    result = await db.partner_orders.find_one_and_update(
        {"id": order_id},
        {"$set": {"status": order_status}},
        return_document=True,
        projection={"_id": 0}
    )
    if not result:
        raise HTTPException(status_code=404, detail="Order not found")
    return result

@api_router.delete("/partner-orders/{order_id}")
async def delete_partner_order(order_id: str, admin: dict = Depends(require_admin)):
    order = await db.partner_orders.find_one({"id": order_id}, {"_id": 0})
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    # Restore partner store debt
    remaining = order.get("remaining", 0)
    if remaining > 0:
        partner = await db.partner_stores.find_one({"id": order["partner_store_id"]})
        if partner:
            new_debt = max(0, partner.get("current_debt", 0) - remaining)
            await db.partner_stores.update_one({"id": order["partner_store_id"]}, {"$set": {"current_debt": new_debt}})
    await db.partner_orders.delete_one({"id": order_id})
    await db.returns.delete_many({"source_id": order_id})
    return {"message": "Pesanan berhasil dihapus"}


@api_router.post("/partner-orders/{order_id}/pay")
async def pay_partner_order(order_id: str, amount: float, current_user: dict = Depends(get_current_user)):
    order = await db.partner_orders.find_one({"id": order_id})
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    
    new_paid = order["paid"] + amount
    new_remaining = order["total"] - new_paid
    
    result = await db.partner_orders.find_one_and_update(
        {"id": order_id},
        {"$set": {"paid": new_paid, "remaining": new_remaining}},
        return_document=True,
        projection={"_id": 0}
    )
    
    # Update partner store debt
    partner_store = await db.partner_stores.find_one({"id": order["partner_store_id"]})
    if partner_store:
        new_debt = max(0, partner_store.get("current_debt", 0) - amount)
        await db.partner_stores.update_one(
            {"id": order["partner_store_id"]},
            {"$set": {"current_debt": new_debt}}
        )
    
    return result


# ============= RETURNS / RETUR =============
class ReturnCreate(BaseModel):
    source_type: str  # "transaction" or "partner_order"
    source_id: str
    items: List[dict]  # [{product_id, product_name, quantity, reason}]
    notes: str = ""
    warehouse_id: str = "gudang_toko"

@api_router.post("/returns")
async def create_return(data: ReturnCreate, admin: dict = Depends(require_admin)):
    now = datetime.now(timezone.utc)
    return_doc = {
        "id": str(uuid.uuid4()),
        "return_number": f"RTN{now.strftime('%Y%m%d%H%M%S')}",
        "source_type": data.source_type,
        "source_id": data.source_id,
        "items": data.items,
        "notes": data.notes,
        "warehouse_id": data.warehouse_id,
        "created_at": now.isoformat(),
        "date": now.strftime("%Y-%m-%d")
    }
    await db.returns.insert_one(return_doc)

    # Restore stock to warehouse
    for item in data.items:
        await db.warehouse_stock.update_one(
            {"product_id": item["product_id"], "warehouse_id": data.warehouse_id},
            {"$inc": {"stock": item["quantity"]}},
            upsert=True
        )
        await recalculate_product_stock(item["product_id"])

    # Mark source as having returns
    if data.source_type == "transaction":
        await db.transactions.update_one({"id": data.source_id}, {"$set": {"has_return": True}})
    elif data.source_type == "partner_order":
        await db.partner_orders.update_one({"id": data.source_id}, {"$set": {"has_return": True}})
        # Reduce partner debt by return total if partner order
        order = await db.partner_orders.find_one({"id": data.source_id}, {"_id": 0})
        if order:
            return_total = sum(i.get("subtotal", 0) for i in data.items)
            new_remaining = max(0, order.get("remaining", 0) - return_total)
            await db.partner_orders.update_one({"id": data.source_id}, {"$set": {"remaining": new_remaining}})
            partner_store = await db.partner_stores.find_one({"id": order["partner_store_id"]})
            if partner_store:
                new_debt = max(0, partner_store.get("current_debt", 0) - return_total)
                await db.partner_stores.update_one({"id": order["partner_store_id"]}, {"$set": {"current_debt": new_debt}})

    return_doc.pop("_id", None)
    return return_doc

@api_router.get("/returns")
async def get_returns(source_id: str = None, admin: dict = Depends(require_admin)):
    query = {}
    if source_id:
        query["source_id"] = source_id
    returns = await db.returns.find(query, {"_id": 0}).sort("created_at", -1).to_list(1000)
    return returns

# ============= CANVASSERS =============
@api_router.get("/canvassers", response_model=List[Canvasser])
async def get_canvassers(current_user: dict = Depends(get_current_user)):
    canvassers = await db.canvassers.find({}, {"_id": 0}).sort("name", 1).to_list(1000)
    return canvassers

@api_router.get("/canvassers/my-orders", response_model=List[PartnerOrder])
async def get_my_orders(current_user: dict = Depends(get_current_user)):
    if current_user.get("role") != "canvasser":
        raise HTTPException(status_code=403, detail="Canvasser access only")
    
    canvasser_id = current_user.get("canvasser_id")
    if not canvasser_id:
        raise HTTPException(status_code=400, detail="Canvasser ID not found")
    
    orders = await db.partner_orders.find(
        {"canvasser_id": canvasser_id},
        {"_id": 0}
    ).sort("created_at", -1).to_list(1000)
    
    return orders

@api_router.get("/canvassers/notifications")
async def get_canvasser_notifications(current_user: dict = Depends(get_current_user)):
    if current_user.get("role") != "canvasser":
        raise HTTPException(status_code=403, detail="Canvasser access only")
    
    canvasser_id = current_user.get("canvasser_id")
    if not canvasser_id:
        return {"new_orders": 0, "pending_orders": 0, "total_remaining": 0}
    
    pending = await db.partner_orders.count_documents({"canvasser_id": canvasser_id, "status": "pending"})
    
    orders = await db.partner_orders.find(
        {"canvasser_id": canvasser_id},
        {"_id": 0, "remaining": 1}
    ).to_list(1000)
    total_remaining = sum(o.get("remaining", 0) for o in orders)
    
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    new_today = await db.partner_orders.count_documents({"canvasser_id": canvasser_id, "date": today})
    
    return {"new_orders": new_today, "pending_orders": pending, "total_remaining": total_remaining}

@api_router.get("/canvassers/product-search")
async def canvasser_product_search(q: str = "", current_user: dict = Depends(get_current_user)):
    if not q or len(q) < 1:
        return []
    query = {"name": {"$regex": q, "$options": "i"}}
    products = await db.products.find(query, {"_id": 0}).sort("name", 1).to_list(50)
    warehouses = await db.warehouses.find({}, {"_id": 0}).to_list(100)
    wh_map = {w["id"]: w["name"] for w in warehouses}
    result = []
    for p in products:
        stocks = await db.warehouse_stock.find({"product_id": p["id"]}, {"_id": 0}).to_list(10)
        wh_stocks = [{"warehouse": wh_map.get(s["warehouse_id"], s["warehouse_id"]), "stock": s.get("stock", 0)} for s in stocks if s.get("stock", 0) > 0]
        result.append({
            "id": p["id"], "name": p["name"], "barcode": p.get("barcode", ""),
            "stock": p.get("stock", 0), "price": p.get("price", 0),
            "box_price": p.get("box_price", 0), "items_per_box": p.get("items_per_box", 1),
            "cost_price": p.get("cost_price", 0),
            "warehouse_stock": wh_stocks
        })
    return result

@api_router.post("/canvassers", response_model=Canvasser)
async def create_canvasser(data: CanvasserCreate, admin: dict = Depends(require_admin)):
    canvasser_doc = {
        "id": str(uuid.uuid4()),
        **data.model_dump(),
        "is_active": True,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.canvassers.insert_one(canvasser_doc)
    return canvasser_doc

@api_router.put("/canvassers/{canvasser_id}", response_model=Canvasser)
async def update_canvasser(canvasser_id: str, data: CanvasserCreate, admin: dict = Depends(require_admin)):
    result = await db.canvassers.find_one_and_update(
        {"id": canvasser_id},
        {"$set": data.model_dump()},
        return_document=True,
        projection={"_id": 0}
    )
    if not result:
        raise HTTPException(status_code=404, detail="Canvasser not found")
    return result

@api_router.delete("/canvassers/{canvasser_id}")
async def delete_canvasser(canvasser_id: str, admin: dict = Depends(require_admin)):
    result = await db.canvassers.delete_one({"id": canvasser_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Canvasser not found")
    return {"message": "Canvasser deleted"}

# ============= PRODUCT ANALYTICS =============
@api_router.get("/analytics/products")
async def get_product_analytics(admin: dict = Depends(require_admin)):
    thirty_days_ago = (datetime.now(timezone.utc) - timedelta(days=30)).strftime("%Y-%m-%d")
    
    transactions = await db.transactions.find(
        {"date": {"$gte": thirty_days_ago}},
        {"_id": 0, "items": 1}
    ).to_list(10000)
    
    product_sales = {}
    for trans in transactions:
        for item in trans.get("items", []):
            pid = item.get("product_id")
            if pid:
                if pid not in product_sales:
                    product_sales[pid] = {"id": pid, "name": item.get("product_name"), "quantity": 0, "revenue": 0}
                product_sales[pid]["quantity"] += item.get("quantity", 0)
                product_sales[pid]["revenue"] += item.get("subtotal", 0)
    
    fast_moving = sorted(product_sales.values(), key=lambda x: x["quantity"], reverse=True)[:10]
    slow_moving = sorted(product_sales.values(), key=lambda x: x["quantity"])[:10]
    
    return {
        "fast_moving": fast_moving,
        "slow_moving": slow_moving
    }

@api_router.get("/analytics/summary")
async def get_analytics_summary(admin: dict = Depends(require_admin)):
    now = datetime.now(timezone.utc)
    this_month_start = now.strftime("%Y-%m-01")
    thirty_days_ago = (now - timedelta(days=30)).strftime("%Y-%m-%d")
    
    # Monthly sales
    month_transactions = await db.transactions.find(
        {"date": {"$gte": this_month_start}}, {"_id": 0}
    ).to_list(10000)
    monthly_sales = sum(t["total"] for t in month_transactions)
    monthly_cost = sum(
        sum(item.get("price", 0) * item.get("quantity", 0) * 0.7 for item in t.get("items", []))
        for t in month_transactions
    )
    
    # Monthly purchases (from suppliers)
    month_purchases = await db.purchases.find(
        {"date": {"$gte": this_month_start}}, {"_id": 0}
    ).to_list(10000)
    monthly_purchases_total = sum(p.get("total", 0) for p in month_purchases)
    monthly_purchases_unpaid = sum(p.get("remaining", 0) for p in month_purchases)
    
    # Partner store stats
    all_partner_orders = await db.partner_orders.find({}, {"_id": 0}).to_list(10000)
    total_partner_sales = sum(o.get("total", 0) for o in all_partner_orders)
    total_partner_unpaid = sum(o.get("remaining", 0) for o in all_partner_orders)
    
    # Daily sales for chart (last 30 days)
    daily_sales = {}
    for i in range(29, -1, -1):
        date = (now - timedelta(days=i)).strftime("%Y-%m-%d")
        daily_sales[date] = {"sales": 0, "transactions": 0}
    
    all_30_transactions = await db.transactions.find(
        {"date": {"$gte": thirty_days_ago}}, {"_id": 0}
    ).to_list(10000)
    for t in all_30_transactions:
        d = t.get("date", "")
        if d in daily_sales:
            daily_sales[d]["sales"] += t.get("total", 0)
            daily_sales[d]["transactions"] += 1
    
    chart_data = [{"date": k, **v} for k, v in daily_sales.items()]
    
    # Category breakdown
    products = await db.products.find({}, {"_id": 0}).to_list(10000)
    categories = await db.categories.find({}, {"_id": 0}).to_list(1000)
    cat_map = {c["id"]: c["name"] for c in categories}
    
    category_stock = {}
    for p in products:
        cat_name = cat_map.get(p.get("category_id", ""), "Tanpa Kategori")
        if cat_name not in category_stock:
            category_stock[cat_name] = {"name": cat_name, "count": 0, "stock_value": 0}
        category_stock[cat_name]["count"] += 1
        category_stock[cat_name]["stock_value"] += p.get("stock", 0) * p.get("cost_price", 0)
    
    return {
        "monthly_sales": monthly_sales,
        "monthly_transactions": len(month_transactions),
        "monthly_purchases": monthly_purchases_total,
        "monthly_purchases_unpaid": monthly_purchases_unpaid,
        "total_partner_sales": total_partner_sales,
        "total_partner_unpaid": total_partner_unpaid,
        "estimated_margin": monthly_sales - monthly_cost,
        "chart_data": chart_data,
        "category_breakdown": list(category_stock.values())
    }

# ============= WAREHOUSES =============
WAREHOUSE_TYPES = {
    "gudang_utama": {"name": "Gudang Utama", "unit": "dus"},
    "gudang_toko": {"name": "Gudang Toko", "unit": "satuan"},
    "gudang_lainnya": {"name": "Gudang Lainnya", "unit": "dus"},
}

@api_router.get("/warehouses")
async def get_warehouses(current_user: dict = Depends(get_current_user)):
    warehouses = await db.warehouses.find({}, {"_id": 0}).sort("sort_order", 1).to_list(100)
    if not warehouses:
        # Seed default warehouses
        defaults = [
            {"id": "gudang_utama", "name": "Gudang Utama", "unit": "dus", "sort_order": 1},
            {"id": "gudang_toko", "name": "Gudang Toko", "unit": "satuan", "sort_order": 2},
            {"id": "gudang_lainnya", "name": "Gudang Lainnya", "unit": "dus", "sort_order": 3},
        ]
        await db.warehouses.insert_many(defaults)
        warehouses = defaults
    return [{k: v for k, v in w.items() if k != "_id"} for w in warehouses]

@api_router.get("/warehouse-stock")
async def get_warehouse_stock(
    product_id: Optional[str] = None,
    warehouse_id: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    query = {}
    if product_id:
        query["product_id"] = product_id
    if warehouse_id:
        query["warehouse_id"] = warehouse_id
    stocks = await db.warehouse_stock.find(query, {"_id": 0}).to_list(10000)
    return stocks

@api_router.put("/warehouse-stock")
async def update_warehouse_stock(
    product_id: str,
    warehouse_id: str,
    stock: int,
    admin: dict = Depends(require_admin)
):
    await db.warehouse_stock.update_one(
        {"product_id": product_id, "warehouse_id": warehouse_id},
        {"$set": {"stock": stock}},
        upsert=True
    )
    # Recalculate total product stock
    await recalculate_product_stock(product_id)
    return {"message": "Stock updated"}

class StockAdjustmentCreate(BaseModel):
    product_id: str
    warehouse_id: str
    new_stock: int
    reason: str = ""

@api_router.post("/stock-adjustments")
async def create_stock_adjustment(
    data: StockAdjustmentCreate,
    admin: dict = Depends(require_admin)
):
    product = await db.products.find_one({"id": data.product_id}, {"_id": 0})
    if not product:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    wh_entry = await db.warehouse_stock.find_one(
        {"product_id": data.product_id, "warehouse_id": data.warehouse_id}
    )
    old_stock = wh_entry["stock"] if wh_entry else 0
    diff = data.new_stock - old_stock
    adjustment = {
        "id": str(uuid.uuid4()),
        "product_id": data.product_id,
        "product_name": product.get("name", ""),
        "warehouse_id": data.warehouse_id,
        "old_stock": old_stock,
        "new_stock": data.new_stock,
        "difference": diff,
        "reason": data.reason,
        "adjusted_by": admin.get("name", admin.get("email", "")),
        "date": datetime.now(timezone.utc).isoformat()
    }
    await db.stock_adjustments.insert_one(adjustment)
    await db.warehouse_stock.update_one(
        {"product_id": data.product_id, "warehouse_id": data.warehouse_id},
        {"$set": {"stock": data.new_stock}},
        upsert=True
    )
    await recalculate_product_stock(data.product_id)
    return {"message": f"Stok disesuaikan: {old_stock} -> {data.new_stock} ({'+' if diff >= 0 else ''}{diff})", "adjustment": {k: v for k, v in adjustment.items() if k != '_id'}}

@api_router.get("/stock-adjustments")
async def get_stock_adjustments(
    product_id: Optional[str] = None,
    admin: dict = Depends(require_admin)
):
    query = {}
    if product_id:
        query["product_id"] = product_id
    adjustments = await db.stock_adjustments.find(query, {"_id": 0}).sort("date", -1).to_list(200)
    return adjustments

async def recalculate_product_stock(product_id: str):
    """Recalculate total product stock in satuan equivalent"""
    product = await db.products.find_one({"id": product_id}, {"_id": 0})
    if not product:
        return
    items_per_box = product.get("items_per_box", 1) or 1
    stocks = await db.warehouse_stock.find({"product_id": product_id}, {"_id": 0}).to_list(100)
    total = 0
    for s in stocks:
        wh = await db.warehouses.find_one({"id": s["warehouse_id"]}, {"_id": 0})
        if wh and wh.get("unit") == "dus":
            total += s["stock"] * items_per_box
        else:
            total += s["stock"]
    await db.products.update_one({"id": product_id}, {"$set": {"stock": total}})

async def ensure_warehouse_stock(product_id: str):
    """Ensure warehouse stock entries exist for all warehouses"""
    warehouses = await db.warehouses.find({}, {"_id": 0}).to_list(100)
    for wh in warehouses:
        exists = await db.warehouse_stock.find_one({"product_id": product_id, "warehouse_id": wh["id"]})
        if not exists:
            await db.warehouse_stock.insert_one({
                "product_id": product_id,
                "warehouse_id": wh["id"],
                "stock": 0
            })

# ============= STOCK TRANSFERS =============
class StockTransferCreate(BaseModel):
    product_id: str
    from_warehouse_id: str
    to_warehouse_id: str
    quantity: int
    notes: Optional[str] = ""

@api_router.get("/stock-transfers")
async def get_stock_transfers(
    product_id: Optional[str] = None,
    admin: dict = Depends(require_admin)
):
    query = {}
    if product_id:
        query["product_id"] = product_id
    transfers = await db.stock_transfers.find(query, {"_id": 0}).sort("created_at", -1).to_list(1000)
    return transfers

@api_router.post("/stock-transfers")
async def create_stock_transfer(data: StockTransferCreate, admin: dict = Depends(require_admin)):
    if data.from_warehouse_id == data.to_warehouse_id:
        raise HTTPException(status_code=400, detail="Gudang asal dan tujuan tidak boleh sama")
    if data.quantity <= 0:
        raise HTTPException(status_code=400, detail="Jumlah harus lebih dari 0")

    product = await db.products.find_one({"id": data.product_id}, {"_id": 0})
    if not product:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")

    from_wh = await db.warehouses.find_one({"id": data.from_warehouse_id}, {"_id": 0})
    to_wh = await db.warehouses.find_one({"id": data.to_warehouse_id}, {"_id": 0})
    if not from_wh or not to_wh:
        raise HTTPException(status_code=404, detail="Gudang tidak ditemukan")

    # Check source stock
    from_stock_doc = await db.warehouse_stock.find_one(
        {"product_id": data.product_id, "warehouse_id": data.from_warehouse_id}
    )
    from_stock = from_stock_doc["stock"] if from_stock_doc else 0
    if from_stock < data.quantity:
        raise HTTPException(status_code=400, detail=f"Stok {from_wh['name']} tidak cukup ({from_stock} {from_wh['unit']})")

    items_per_box = product.get("items_per_box", 1) or 1
    from_unit = from_wh["unit"]
    to_unit = to_wh["unit"]

    # Calculate converted quantity
    if from_unit == "dus" and to_unit == "satuan":
        converted_qty = data.quantity * items_per_box
    elif from_unit == "satuan" and to_unit == "dus":
        if data.quantity % items_per_box != 0:
            raise HTTPException(
                status_code=400,
                detail=f"Jumlah satuan ({data.quantity}) harus kelipatan {items_per_box} untuk konversi ke dus"
            )
        converted_qty = data.quantity // items_per_box
    else:
        converted_qty = data.quantity

    # Deduct from source
    await db.warehouse_stock.update_one(
        {"product_id": data.product_id, "warehouse_id": data.from_warehouse_id},
        {"$inc": {"stock": -data.quantity}}
    )
    # Add to destination
    await db.warehouse_stock.update_one(
        {"product_id": data.product_id, "warehouse_id": data.to_warehouse_id},
        {"$inc": {"stock": converted_qty}},
        upsert=True
    )

    now = datetime.now(timezone.utc)
    transfer_doc = {
        "id": str(uuid.uuid4()),
        "product_id": data.product_id,
        "product_name": product["name"],
        "from_warehouse_id": data.from_warehouse_id,
        "from_warehouse_name": from_wh["name"],
        "to_warehouse_id": data.to_warehouse_id,
        "to_warehouse_name": to_wh["name"],
        "quantity": data.quantity,
        "from_unit": from_unit,
        "converted_quantity": converted_qty,
        "to_unit": to_unit,
        "notes": data.notes,
        "created_at": now.isoformat(),
        "date": now.strftime("%Y-%m-%d")
    }
    await db.stock_transfers.insert_one(transfer_doc)

    await recalculate_product_stock(data.product_id)

    return {k: v for k, v in transfer_doc.items() if k != "_id"}

# ============= PARTNER PRICING =============
@api_router.get("/partner-prices")
async def get_partner_prices(
    partner_store_id: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    query = {}
    if partner_store_id:
        query["partner_store_id"] = partner_store_id
    prices = await db.partner_prices.find(query, {"_id": 0}).to_list(10000)
    return prices

@api_router.put("/partner-prices")
async def set_partner_price(
    partner_store_id: str,
    product_id: str,
    custom_price: float,
    admin: dict = Depends(require_admin)
):
    await db.partner_prices.update_one(
        {"partner_store_id": partner_store_id, "product_id": product_id},
        {"$set": {"custom_price": custom_price, "updated_at": datetime.now(timezone.utc).isoformat()}},
        upsert=True
    )
    return {"message": "Harga mitra diperbarui"}

@api_router.delete("/partner-prices")
async def delete_partner_price(
    partner_store_id: str,
    product_id: str,
    admin: dict = Depends(require_admin)
):
    await db.partner_prices.delete_one(
        {"partner_store_id": partner_store_id, "product_id": product_id}
    )
    return {"message": "Harga mitra dihapus"}

# ============= SUPPLIER PURCHASE DETAILS =============
@api_router.get("/suppliers/{supplier_id}/purchases")
async def get_supplier_purchases(supplier_id: str, admin: dict = Depends(require_admin)):
    purchases = await db.purchases.find({"supplier_id": supplier_id}, {"_id": 0}).sort("created_at", -1).to_list(1000)
    return purchases

@api_router.get("/products/{product_id}/price-comparison")
async def get_product_price_comparison(product_id: str, admin: dict = Depends(require_admin)):
    """Get all purchase prices for a product from different suppliers"""
    all_purchases = await db.purchases.find({}, {"_id": 0}).to_list(10000)
    price_data = {}
    for purchase in all_purchases:
        for item in purchase.get("items", []):
            if item.get("product_id") == product_id:
                sid = purchase["supplier_id"]
                sname = purchase["supplier_name"]
                if sid not in price_data:
                    price_data[sid] = {"supplier_id": sid, "supplier_name": sname, "prices": [], "last_price": 0, "total_qty": 0}
                price_data[sid]["prices"].append({"price": item.get("cost_price", 0), "date": purchase.get("date", ""), "quantity": item.get("quantity", 0)})
                price_data[sid]["last_price"] = item.get("cost_price", 0)
                price_data[sid]["total_qty"] += item.get("quantity", 0)
    return list(price_data.values())

# ============= VEHICLES (ARMADA MOBIL BOX) =============
class VehicleCreate(BaseModel):
    name: str
    plate_number: str
    area_route: Optional[str] = ""
    driver_name: Optional[str] = ""
    canvasser_id: Optional[str] = ""

class Vehicle(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    name: str
    plate_number: str
    area_route: Optional[str] = ""
    driver_name: Optional[str] = ""
    canvasser_id: Optional[str] = ""
    created_at: str

@api_router.get("/vehicles")
async def get_vehicles(current_user: dict = Depends(get_current_user)):
    vehicles = await db.vehicles.find({}, {"_id": 0}).sort("name", 1).to_list(100)
    return vehicles

@api_router.post("/vehicles")
async def create_vehicle(data: VehicleCreate, admin: dict = Depends(require_admin_or_manager)):
    doc = {"id": str(uuid.uuid4()), **data.model_dump(), "created_at": datetime.now(timezone.utc).isoformat()}
    await db.vehicles.insert_one(doc)
    await log_audit(admin, "create", "vehicle", doc["id"], {"name": data.name, "plate": data.plate_number})
    return {k: v for k, v in doc.items() if k != '_id'}

@api_router.put("/vehicles/{vehicle_id}")
async def update_vehicle(vehicle_id: str, data: VehicleCreate, admin: dict = Depends(require_admin_or_manager)):
    result = await db.vehicles.find_one_and_update(
        {"id": vehicle_id}, {"$set": data.model_dump()}, return_document=True, projection={"_id": 0}
    )
    if not result:
        raise HTTPException(status_code=404, detail="Armada tidak ditemukan")
    await log_audit(admin, "update", "vehicle", vehicle_id, data.model_dump())
    return result

@api_router.delete("/vehicles/{vehicle_id}")
async def delete_vehicle(vehicle_id: str, admin: dict = Depends(require_admin_or_manager)):
    result = await db.vehicles.delete_one({"id": vehicle_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Armada tidak ditemukan")
    await db.partner_stores.update_many({"vehicle_id": vehicle_id}, {"$set": {"vehicle_id": ""}})
    await log_audit(admin, "delete", "vehicle", vehicle_id, {})
    return {"message": "Armada berhasil dihapus"}

# ============= ROLE PERMISSIONS =============
@api_router.get("/role-permissions")
async def get_role_permissions(current_user: dict = Depends(get_current_user)):
    return {"permissions": ROLE_PERMISSIONS, "current_role": current_user.get("role", ""), "current_permissions": ROLE_PERMISSIONS.get(current_user.get("role", ""), [])}

# ============= AUDIT TRAIL =============
async def log_audit(user: dict, action: str, entity_type: str, entity_id: str, details: dict):
    entry = {
        "id": str(uuid.uuid4()),
        "user_id": user.get("id", ""),
        "user_name": user.get("name", user.get("email", "")),
        "user_role": user.get("role", ""),
        "action": action,
        "entity_type": entity_type,
        "entity_id": entity_id,
        "details": details,
        "timestamp": datetime.now(timezone.utc).isoformat()
    }
    await db.audit_log.insert_one(entry)

@api_router.get("/audit-log")
async def get_audit_log(
    entity_type: Optional[str] = None,
    limit: int = 100,
    admin: dict = Depends(require_admin_or_manager)
):
    query = {}
    if entity_type:
        query["entity_type"] = entity_type
    logs = await db.audit_log.find(query, {"_id": 0}).sort("timestamp", -1).to_list(limit)
    return logs

# ============= DATA IMPORT (Bypass Business Logic) =============
IMPORT_TYPES_INFO = [
    {"key": "warehouses", "label": "Gudang", "order": 1, "columns": ["nama", "unit", "urutan"]},
    {"key": "categories", "label": "Kategori", "order": 2, "columns": ["nama", "deskripsi"]},
    {"key": "suppliers", "label": "Supplier", "order": 3, "columns": ["nama", "kontak", "telepon", "email", "alamat"]},
    {"key": "products", "label": "Produk", "order": 4, "columns": ["nama", "barcode", "kategori", "supplier", "harga_jual", "harga_beli", "harga_box", "isi_per_box", "min_stok", "kadaluarsa", "batch", "deskripsi"]},
    {"key": "warehouse_stock", "label": "Stok per Gudang", "order": 5, "columns": ["produk", "gudang", "stok"]},
    {"key": "purchases", "label": "Pembelian Supplier", "order": 6, "columns": ["no_pembelian", "supplier", "tanggal", "produk", "jumlah", "harga_satuan", "total_pembelian", "dibayar", "sisa", "gudang", "catatan"]},
    {"key": "payment_history", "label": "Pembayaran Hutang Supplier", "order": 7, "columns": ["supplier", "tanggal", "jumlah", "catatan"]},
    {"key": "partner_stores", "label": "Kios Mitra", "order": 8, "columns": ["nama", "pemilik", "telepon", "alamat", "limit_kredit", "hutang_saat_ini"]},
    {"key": "partner_orders", "label": "Pesanan Kios Mitra", "order": 9, "columns": ["no_pesanan", "kios_mitra", "tanggal", "produk", "jumlah", "harga", "gudang", "diskon", "total_pesanan", "dibayar", "sisa", "sales", "status"]},
]

async def parse_upload(file: UploadFile):
    content = await file.read()
    fname = (file.filename or "").lower()
    if fname.endswith('.xlsx') or fname.endswith('.xls'):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
        headers = []
        for cell in ws[1]:
            if cell.value is not None:
                headers.append(str(cell.value).strip().lower())
            else:
                break
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            rd = {}
            for idx, val in enumerate(row):
                if idx < len(headers):
                    rd[headers[idx]] = val if val is not None else ""
            if any(str(v).strip() != "" for v in rd.values()):
                rows.append(rd)
        return rows
    else:
        text = content.decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(text))
        return [{k.strip().lower(): (v.strip() if isinstance(v, str) else v) for k, v in row.items() if k} for row in reader]

def safe_float(v, default=0):
    try:
        return float(str(v).replace(",", ".").replace(" ", "")) if v != "" and v is not None else default
    except Exception:
        return default

def safe_int(v, default=0):
    try:
        return int(float(str(v).replace(",", ".").replace(" ", ""))) if v != "" and v is not None else default
    except Exception:
        return default

def safe_str(v, default=""):
    return str(v).strip() if v is not None and str(v).strip() else default

@api_router.get("/import/types")
async def get_import_types(admin: dict = Depends(require_admin)):
    return IMPORT_TYPES_INFO

@api_router.get("/import/counts")
async def get_import_counts(admin: dict = Depends(require_admin)):
    collections = ["warehouses", "categories", "suppliers", "products", "warehouse_stock", "purchases", "payment_history", "partner_stores", "partner_orders"]
    result = {}
    for c in collections:
        result[c] = await db[c].count_documents({})
    return result

@api_router.get("/import/template/{import_type}")
async def get_import_template(import_type: str, admin: dict = Depends(require_admin)):
    info = next((t for t in IMPORT_TYPES_INFO if t["key"] == import_type), None)
    if not info:
        raise HTTPException(status_code=404, detail="Tipe import tidak ditemukan")
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(info["columns"])
    output.seek(0)
    from starlette.responses import StreamingResponse
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=template_{import_type}.csv"}
    )

@api_router.post("/import/{import_type}")
async def import_data(import_type: str, file: UploadFile = File(...), admin: dict = Depends(require_admin)):
    try:
        rows = await parse_upload(file)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Gagal membaca file: {str(e)}")
    if not rows:
        raise HTTPException(status_code=400, detail="File kosong atau format tidak sesuai")

    now = datetime.now(timezone.utc)
    results = {"success": 0, "skipped": 0, "errors": []}

    if import_type == "warehouses":
        existing = {w["name"].lower(): w for w in await db.warehouses.find({}, {"_id": 0}).to_list(100)}
        max_order = max((w.get("sort_order", 0) for w in existing.values()), default=0)
        for i, row in enumerate(rows, 1):
            nama = safe_str(row.get("nama"))
            if not nama:
                results["errors"].append(f"Baris {i}: Nama kosong")
                continue
            if nama.lower() in existing:
                results["skipped"] += 1
                continue
            max_order += 1
            doc = {"id": nama.lower().replace(" ", "_").replace("-", "_"), "name": nama, "unit": safe_str(row.get("unit"), "dus").lower(), "sort_order": safe_int(row.get("urutan"), max_order)}
            if doc["unit"] not in ("dus", "satuan"):
                doc["unit"] = "dus"
            await db.warehouses.insert_one(doc)
            existing[nama.lower()] = doc
            results["success"] += 1

    elif import_type == "categories":
        existing = {c["name"].lower() for c in await db.categories.find({}, {"_id": 0, "name": 1}).to_list(10000)}
        for i, row in enumerate(rows, 1):
            nama = safe_str(row.get("nama"))
            if not nama:
                results["errors"].append(f"Baris {i}: Nama kosong")
                continue
            if nama.lower() in existing:
                results["skipped"] += 1
                continue
            doc = {"id": str(uuid.uuid4()), "name": nama, "description": safe_str(row.get("deskripsi")), "created_at": now.isoformat()}
            await db.categories.insert_one(doc)
            existing.add(nama.lower())
            results["success"] += 1

    elif import_type == "suppliers":
        existing = {s["name"].lower() for s in await db.suppliers.find({}, {"_id": 0, "name": 1}).to_list(10000)}
        for i, row in enumerate(rows, 1):
            nama = safe_str(row.get("nama"))
            if not nama:
                results["errors"].append(f"Baris {i}: Nama kosong")
                continue
            if nama.lower() in existing:
                results["skipped"] += 1
                continue
            doc = {"id": str(uuid.uuid4()), "name": nama, "contact_person": safe_str(row.get("kontak")), "phone": safe_str(row.get("telepon")), "email": safe_str(row.get("email")), "address": safe_str(row.get("alamat")), "created_at": now.isoformat()}
            await db.suppliers.insert_one(doc)
            existing.add(nama.lower())
            results["success"] += 1

    elif import_type == "products":
        cats = {c["name"].lower(): c["id"] for c in await db.categories.find({}, {"_id": 0}).to_list(10000)}
        sups = {s["name"].lower(): s["id"] for s in await db.suppliers.find({}, {"_id": 0}).to_list(10000)}
        existing = {p["name"].lower() for p in await db.products.find({}, {"_id": 0, "name": 1}).to_list(100000)}
        for i, row in enumerate(rows, 1):
            nama = safe_str(row.get("nama"))
            if not nama:
                results["errors"].append(f"Baris {i}: Nama kosong")
                continue
            if nama.lower() in existing:
                results["skipped"] += 1
                continue
            cat_name = safe_str(row.get("kategori")).lower()
            sup_name = safe_str(row.get("supplier")).lower()
            cat_id = cats.get(cat_name, "")
            sup_id = sups.get(sup_name, "")
            if cat_name and not cat_id:
                results["errors"].append(f"Baris {i}: Kategori '{row.get('kategori')}' tidak ditemukan")
                continue
            if sup_name and not sup_id:
                results["errors"].append(f"Baris {i}: Supplier '{row.get('supplier')}' tidak ditemukan")
                continue
            doc = {
                "id": str(uuid.uuid4()), "name": nama, "barcode": safe_str(row.get("barcode")),
                "category_id": cat_id, "supplier_id": sup_id, "stock": 0,
                "price": safe_float(row.get("harga_jual")), "cost_price": safe_float(row.get("harga_beli")),
                "box_price": safe_float(row.get("harga_box")), "items_per_box": max(1, safe_int(row.get("isi_per_box"), 1)),
                "min_stock_alert": safe_int(row.get("min_stok"), 10), "expiry_date": safe_str(row.get("kadaluarsa")),
                "batch_number": safe_str(row.get("batch")), "description": safe_str(row.get("deskripsi")),
                "image": "", "created_at": now.isoformat(), "updated_at": now.isoformat()
            }
            await db.products.insert_one(doc)
            existing.add(nama.lower())
            results["success"] += 1

    elif import_type == "warehouse_stock":
        prods = {p["name"].lower(): p for p in await db.products.find({}, {"_id": 0}).to_list(100000)}
        whs = {w["name"].lower(): w["id"] for w in await db.warehouses.find({}, {"_id": 0}).to_list(100)}
        affected = set()
        for i, row in enumerate(rows, 1):
            prod_name = safe_str(row.get("produk")).lower()
            wh_name = safe_str(row.get("gudang")).lower()
            stok = safe_int(row.get("stok"))
            if not prod_name or not wh_name:
                results["errors"].append(f"Baris {i}: Produk atau gudang kosong")
                continue
            prod = prods.get(prod_name)
            wh_id = whs.get(wh_name)
            if not prod:
                results["errors"].append(f"Baris {i}: Produk '{row.get('produk')}' tidak ditemukan")
                continue
            if not wh_id:
                results["errors"].append(f"Baris {i}: Gudang '{row.get('gudang')}' tidak ditemukan")
                continue
            await db.warehouse_stock.update_one({"product_id": prod["id"], "warehouse_id": wh_id}, {"$set": {"stock": stok}}, upsert=True)
            affected.add(prod["id"])
            results["success"] += 1
        for pid in affected:
            await recalculate_product_stock(pid)

    elif import_type == "purchases":
        sups = {s["name"].lower(): s for s in await db.suppliers.find({}, {"_id": 0}).to_list(10000)}
        prods = {p["name"].lower(): p for p in await db.products.find({}, {"_id": 0}).to_list(100000)}
        whs = {w["name"].lower(): w["id"] for w in await db.warehouses.find({}, {"_id": 0}).to_list(100)}
        groups = {}
        for i, row in enumerate(rows, 1):
            key = safe_str(row.get("no_pembelian")) or f"auto_{i}"
            if key not in groups:
                groups[key] = {"rows": [], "indices": []}
            groups[key]["rows"].append(row)
            groups[key]["indices"].append(i)
        for key, group in groups.items():
            first = group["rows"][0]
            sup_name = safe_str(first.get("supplier")).lower()
            sup = sups.get(sup_name)
            if not sup:
                for idx in group["indices"]:
                    results["errors"].append(f"Baris {idx}: Supplier '{first.get('supplier')}' tidak ditemukan")
                continue
            items = []
            has_error = False
            for j, row in enumerate(group["rows"]):
                prod_name = safe_str(row.get("produk")).lower()
                if not prod_name:
                    continue
                prod = prods.get(prod_name)
                if not prod:
                    results["errors"].append(f"Baris {group['indices'][j]}: Produk '{row.get('produk')}' tidak ditemukan")
                    has_error = True
                    continue
                jumlah = safe_int(row.get("jumlah"))
                harga = safe_float(row.get("harga_satuan"))
                items.append({"product_id": prod["id"], "product_name": prod["name"], "quantity": jumlah, "cost_price": harga, "subtotal": jumlah * harga})
            if has_error or not items:
                continue
            total = safe_float(first.get("total_pembelian")) or sum(it["subtotal"] for it in items)
            dibayar = safe_float(first.get("dibayar"))
            sisa = safe_float(first.get("sisa")) if safe_str(first.get("sisa")) else total - dibayar
            tanggal = safe_str(first.get("tanggal")) or now.strftime("%Y-%m-%d")
            gudang_id = whs.get(safe_str(first.get("gudang")).lower(), "gudang_utama")
            doc = {
                "id": str(uuid.uuid4()), "supplier_id": sup["id"], "supplier_name": sup["name"],
                "items": items, "total": total, "paid": dibayar, "remaining": sisa,
                "warehouse_id": gudang_id, "notes": safe_str(first.get("catatan")),
                "created_at": tanggal + "T00:00:00+00:00", "date": tanggal
            }
            await db.purchases.insert_one(doc)
            results["success"] += 1

    elif import_type == "payment_history":
        sups = {s["name"].lower(): s for s in await db.suppliers.find({}, {"_id": 0}).to_list(10000)}
        for i, row in enumerate(rows, 1):
            sup_name = safe_str(row.get("supplier")).lower()
            sup = sups.get(sup_name)
            if not sup:
                results["errors"].append(f"Baris {i}: Supplier '{row.get('supplier')}' tidak ditemukan")
                continue
            jumlah = safe_float(row.get("jumlah"))
            if jumlah <= 0:
                results["errors"].append(f"Baris {i}: Jumlah harus > 0")
                continue
            tanggal = safe_str(row.get("tanggal")) or now.strftime("%Y-%m-%d")
            doc = {
                "id": str(uuid.uuid4()), "type": "supplier", "reference_id": sup["id"],
                "supplier_id": sup["id"], "supplier_name": sup["name"], "amount": jumlah,
                "notes": safe_str(row.get("catatan")) or f"Import: Pembayaran ke {sup['name']}",
                "created_at": tanggal + "T00:00:00+00:00", "date": tanggal
            }
            await db.payment_history.insert_one(doc)
            results["success"] += 1

    elif import_type == "partner_stores":
        existing = {s["name"].lower() for s in await db.partner_stores.find({}, {"_id": 0, "name": 1}).to_list(10000)}
        for i, row in enumerate(rows, 1):
            nama = safe_str(row.get("nama"))
            if not nama:
                results["errors"].append(f"Baris {i}: Nama kosong")
                continue
            if nama.lower() in existing:
                results["skipped"] += 1
                continue
            doc = {
                "id": str(uuid.uuid4()), "name": nama, "owner_name": safe_str(row.get("pemilik")),
                "phone": safe_str(row.get("telepon")), "address": safe_str(row.get("alamat")),
                "latitude": None, "longitude": None,
                "credit_limit": safe_float(row.get("limit_kredit")),
                "current_debt": safe_float(row.get("hutang_saat_ini")),
                "created_at": now.isoformat()
            }
            await db.partner_stores.insert_one(doc)
            existing.add(nama.lower())
            results["success"] += 1

    elif import_type == "partner_orders":
        stores = {s["name"].lower(): s for s in await db.partner_stores.find({}, {"_id": 0}).to_list(10000)}
        prods = {p["name"].lower(): p for p in await db.products.find({}, {"_id": 0}).to_list(100000)}
        whs = {w["name"].lower(): w["id"] for w in await db.warehouses.find({}, {"_id": 0}).to_list(100)}
        canvassers_map = {c["name"].lower(): c for c in await db.canvassers.find({}, {"_id": 0}).to_list(1000)}
        groups = {}
        for i, row in enumerate(rows, 1):
            key = safe_str(row.get("no_pesanan")) or f"auto_{i}"
            if key not in groups:
                groups[key] = {"rows": [], "indices": []}
            groups[key]["rows"].append(row)
            groups[key]["indices"].append(i)
        for key, group in groups.items():
            first = group["rows"][0]
            store_name = safe_str(first.get("kios_mitra")).lower()
            store = stores.get(store_name)
            if not store:
                for idx in group["indices"]:
                    results["errors"].append(f"Baris {idx}: Kios '{first.get('kios_mitra')}' tidak ditemukan")
                continue
            items = []
            has_error = False
            for j, row in enumerate(group["rows"]):
                prod_name = safe_str(row.get("produk")).lower()
                if not prod_name:
                    continue
                prod = prods.get(prod_name)
                if not prod:
                    results["errors"].append(f"Baris {group['indices'][j]}: Produk '{row.get('produk')}' tidak ditemukan")
                    has_error = True
                    continue
                jumlah = safe_int(row.get("jumlah"))
                harga = safe_float(row.get("harga"))
                gudang_id = whs.get(safe_str(row.get("gudang")).lower(), "")
                items.append({"product_id": prod["id"], "product_name": prod["name"], "quantity": jumlah, "price": harga, "subtotal": jumlah * harga, "warehouse_id": gudang_id, "sell_unit": ""})
            if has_error or not items:
                continue
            subtotal = sum(it["subtotal"] for it in items)
            diskon = safe_float(first.get("diskon"))
            total = safe_float(first.get("total_pesanan")) or (subtotal - diskon)
            dibayar = safe_float(first.get("dibayar"))
            sisa = safe_float(first.get("sisa")) if safe_str(first.get("sisa")) else total - dibayar
            tanggal = safe_str(first.get("tanggal")) or now.strftime("%Y-%m-%d")
            sales_name = safe_str(first.get("sales")).lower()
            canvasser = canvassers_map.get(sales_name)
            order_number = key if not key.startswith("auto_") else f"PO-IMP-{str(uuid.uuid4())[:8].upper()}"
            doc = {
                "id": str(uuid.uuid4()), "order_number": order_number,
                "partner_store_id": store["id"], "partner_store_name": store["name"],
                "items": items, "subtotal": subtotal, "discount": diskon,
                "total": total, "paid": dibayar, "remaining": sisa,
                "delivery_date": tanggal,
                "canvasser_id": canvasser["id"] if canvasser else "",
                "canvasser_name": canvasser["name"] if canvasser else safe_str(first.get("sales")),
                "warehouse_id": whs.get(safe_str(first.get("gudang")).lower(), ""),
                "status": safe_str(first.get("status")) or "completed",
                "created_at": tanggal + "T00:00:00+00:00", "date": tanggal, "has_return": False
            }
            await db.partner_orders.insert_one(doc)
            results["success"] += 1
    else:
        raise HTTPException(status_code=400, detail=f"Tipe import '{import_type}' tidak didukung")

    return {"type": import_type, "total_rows": len(rows), "success": results["success"], "skipped": results["skipped"], "error_count": len(results["errors"]), "errors": results["errors"][:50]}

@api_router.delete("/import/clear/{collection_type}")
async def clear_collection(collection_type: str, admin: dict = Depends(require_admin)):
    allowed = ["warehouses", "categories", "suppliers", "products", "warehouse_stock", "purchases", "payment_history", "partner_stores", "partner_orders"]
    if collection_type not in allowed:
        raise HTTPException(status_code=400, detail="Tipe tidak didukung")
    count = await db[collection_type].count_documents({})
    await db[collection_type].delete_many({})
    return {"message": f"{count} data dihapus dari {collection_type}"}

app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

@app.on_event("startup")
async def startup_db():
    admin_exists = await db.users.find_one({"role": "admin"})
    if not admin_exists:
        default_admin = {
            "id": str(uuid.uuid4()),
            "email": "admin@aqintani.com",
            "password": hash_password("admin123"),
            "name": "Administrator",
            "role": "admin",
            "created_at": datetime.now(timezone.utc).isoformat()
        }
        await db.users.insert_one(default_admin)
        logger.info("Default admin created: admin@aqintani.com / admin123")
    
    # Seed warehouses
    wh_count = await db.warehouses.count_documents({})
    if wh_count == 0:
        await db.warehouses.insert_many([
            {"id": "gudang_utama", "name": "Gudang Utama", "unit": "dus", "sort_order": 1},
            {"id": "gudang_toko", "name": "Gudang Toko", "unit": "satuan", "sort_order": 2},
            {"id": "gudang_lainnya", "name": "Gudang Lainnya", "unit": "dus", "sort_order": 3},
        ])
        logger.info("Default warehouses created")
    
    # Migrate existing products: ensure warehouse_stock entries + items_per_box
    products = await db.products.find({}, {"_id": 0}).to_list(10000)
    product_ids = set()
    for p in products:
        product_ids.add(p["id"])
        await ensure_warehouse_stock(p["id"])
        # If product has stock but no warehouse_stock anywhere, put it in Gudang Toko
        all_wh_stock = await db.warehouse_stock.find({"product_id": p["id"]}, {"_id": 0}).to_list(10)
        total_wh = sum(s.get("stock", 0) for s in all_wh_stock)
        if total_wh == 0 and p.get("stock", 0) > 0:
            await db.warehouse_stock.update_one(
                {"product_id": p["id"], "warehouse_id": "gudang_toko"},
                {"$set": {"stock": p["stock"]}}
            )
        # Ensure items_per_box field exists
        if "items_per_box" not in p:
            await db.products.update_one({"id": p["id"]}, {"$set": {"items_per_box": 1, "box_price": 0}})
        # Recalculate total stock from warehouse data
        await recalculate_product_stock(p["id"])
    
    # Clean orphaned warehouse_stock entries (products that were deleted)
    all_wh_stocks = await db.warehouse_stock.find({}, {"_id": 0}).to_list(100000)
    for s in all_wh_stocks:
        if s.get("product_id") not in product_ids:
            await db.warehouse_stock.delete_many({"product_id": s["product_id"]})

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()